"""Вес и объём для КП: факты прайса с приоритетом проверенного справочника ППО."""

from __future__ import annotations

import csv
import io
import re
import uuid
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.matching.text import normalize_article, normalize_name
from app.models import (
    AttributeDefinition,
    ProductPhysicalReference,
    SupplierPriceItem,
    SupplierPriceItemAttributeValue,
)


@dataclass(frozen=True, slots=True)
class PhysicalValues:
    weight_kg: Decimal | None = None
    volume_m3: Decimal | None = None
    weight_source: str = ""
    volume_source: str = ""


ALIASES = {
    "name": {"наименование", "товар", "название", "name", "product"},
    "article": {"артикул", "код", "article", "sku"},
    "weight": {"вес", "вес кг", "масса", "масса кг", "weight", "weight kg"},
    "volume": {"объем", "объем м3", "объём", "объём м3", "volume", "volume m3"},
}


def _header(value: object) -> str:
    raw = str(value or "").casefold().replace("³", "3")
    return " ".join(re.sub(r"[^0-9a-zа-яё]+", " ", raw).split())


def _decimal(value: object) -> Decimal | None:
    raw = str(value or "").strip().replace("\u00a0", "").replace(" ", "").replace(",", ".")
    raw = re.sub(r"[^0-9.+-].*$", "", raw)
    if not raw:
        return None
    try:
        result = Decimal(raw)
    except InvalidOperation:
        return None
    return result if result >= 0 else None


def _rows(payload: bytes, filename: str) -> list[list[object]]:
    if filename.casefold().endswith(".xlsx"):
        workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
        try:
            return [list(row) for row in workbook.active.iter_rows(values_only=True)]
        finally:
            workbook.close()
    text = payload.decode("utf-8-sig", errors="replace")
    first = text.splitlines()[0] if text.splitlines() else ""
    delimiter = max((";", ",", "\t"), key=first.count)
    return [list(row) for row in csv.reader(io.StringIO(text), delimiter=delimiter)]


async def import_reference(
    session: AsyncSession, *, project_id: uuid.UUID, payload: bytes, filename: str
) -> dict[str, int]:
    rows = _rows(payload, filename)
    if not rows:
        raise ValueError("файл справочника пуст")
    indexes: dict[str, int] = {}
    for index, cell in enumerate(rows[0]):
        normalized = _header(cell)
        for field, aliases in ALIASES.items():
            if normalized in aliases:
                indexes[field] = index
    if "name" not in indexes and "article" not in indexes:
        raise ValueError("нужна колонка «Наименование» или «Артикул»")
    if "weight" not in indexes and "volume" not in indexes:
        raise ValueError("нужна колонка «Вес, кг» или «Объём, м3»")

    imported = skipped = 0
    for source_row in rows[1:]:
        def cell(field: str) -> object:
            index = indexes.get(field)
            return source_row[index] if index is not None and index < len(source_row) else ""

        name = str(cell("name") or "").strip()
        article = str(cell("article") or "").strip()
        article_key = normalize_article(article) or None
        name_key = normalize_name(name)
        weight = _decimal(cell("weight"))
        volume = _decimal(cell("volume"))
        if (not article_key and not name_key) or (weight is None and volume is None):
            skipped += 1
            continue
        match_key = f"article:{article_key}" if article_key else f"name:{name_key}"
        existing = await session.scalar(
            select(ProductPhysicalReference).where(
                ProductPhysicalReference.project_id == project_id,
                ProductPhysicalReference.match_key == match_key,
            )
        )
        if existing is None:
            existing = ProductPhysicalReference(project_id=project_id, match_key=match_key)
            session.add(existing)
        existing.normalized_article = article_key
        existing.normalized_name = name_key
        existing.name = name or article
        existing.weight_kg = weight
        existing.volume_m3 = volume
        existing.source = filename
        imported += 1
    await session.flush()
    return {"imported": imported, "skipped": skipped}


async def resolve_physical(
    session: AsyncSession, *, item: SupplierPriceItem, project_id: uuid.UUID
) -> PhysicalValues:
    facts = list(
        await session.execute(
            select(AttributeDefinition.code, SupplierPriceItemAttributeValue.value_decimal)
            .join(
                SupplierPriceItemAttributeValue,
                SupplierPriceItemAttributeValue.attribute_definition_id == AttributeDefinition.id,
            )
            .where(
                SupplierPriceItemAttributeValue.item_id == item.id,
                SupplierPriceItemAttributeValue.parse_status == "parsed",
                SupplierPriceItemAttributeValue.value_decimal.is_not(None),
                AttributeDefinition.code.in_(("weight", "volume")),
            )
        )
    )
    by_code = {code: Decimal(value) for code, value in facts}
    weight = by_code.get("weight")
    volume = by_code.get("volume")
    weight_source = "price" if weight is not None else ""
    volume_source = "price" if volume is not None else ""

    article_key = normalize_article(item.source_article)
    name_key = normalize_name(item.source_name)
    statement = select(ProductPhysicalReference).where(
        ProductPhysicalReference.project_id == project_id
    )
    reference = None
    if article_key:
        reference = await session.scalar(statement.where(
            ProductPhysicalReference.normalized_article == article_key
        ))
    if reference is None and name_key:
        reference = await session.scalar(statement.where(
            ProductPhysicalReference.normalized_name == name_key
        ))
    # ППО — управляемый справочник, поэтому заполненное значение из него
    # перекрывает неоднозначную колонку поставщика. Пропуски дополняются прайсом.
    if reference is not None:
        if reference.weight_kg is not None:
            weight, weight_source = Decimal(reference.weight_kg), "reference"
        if reference.volume_m3 is not None:
            volume, volume_source = Decimal(reference.volume_m3), "reference"
    return PhysicalValues(weight, volume, weight_source, volume_source)
