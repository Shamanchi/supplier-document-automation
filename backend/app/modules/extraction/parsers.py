"""Native parsers and deterministic technical-parameter extraction."""

from __future__ import annotations

import re
import unicodedata
from dataclasses import dataclass
from io import BytesIO

from docx import Document as DocxDocument
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from pypdf import PdfReader

from app.modules.documents.signatures import (
    DOCX_MEDIA_TYPE,
    PDF_MEDIA_TYPE,
    XLSX_MEDIA_TYPE,
)
from app.modules.extraction.table_detection import (
    detect_tables_from_grid,
    detect_tables_on_page,
    group_words_into_lines,
    merge_page_tables,
)

PARSER_VERSION = "native-v1"
MAX_PAGES = 500
MAX_ROWS_PER_SHEET = 10_000
MAX_FRAGMENT_CHARS = 20_000

_PAIR = re.compile(r"^([^:=]{2,160}?)\s*[:=]\s*(\S.+)$")
_NUMBER_WITH_UNIT = re.compile(
    r"^([<>]=?\s*)?([-+]?\d+(?:[.,]\d+)?(?:\s*(?:-|–|…|\.\.)\s*"
    r"[-+]?\d+(?:[.,]\d+)?)?)\s*([^\d\s].*)?$"
)
_HEADER_WORDS = {
    "параметр",
    "характеристика",
    "наименование",
    "показатель",
    "значение",
    "единица",
    "ед. изм.",
    "unit",
    "value",
    "parameter",
}


class NativeParseError(ValueError):
    """The document cannot be parsed safely by the native contour."""


@dataclass(frozen=True)
class ParsedFragment:
    text: str
    method: str
    locator_kind: str
    locator: dict
    cells: tuple[str, ...] | None = None
    cell_boxes: tuple[dict, ...] | None = None
    table_bbox: tuple[float | None, ...] | None = None
    table_source: str | None = None
    parser_version: str = PARSER_VERSION
    confidence: float | None = None


@dataclass(frozen=True)
class ParsedPage:
    page_no: int
    fragments: tuple[ParsedFragment, ...]
    quality_score: float
    ocr_needed: bool
    document_type: str = "NATIVE"
    active_layer: str = "NATIVE"
    dpi: int | None = None
    width_px: int | None = None
    height_px: int | None = None
    parser_version: str = PARSER_VERSION


@dataclass(frozen=True)
class ParsedDocument:
    pages: tuple[ParsedPage, ...]
    page_words: tuple[tuple[tuple, ...], ...] = ()


@dataclass(frozen=True)
class ParameterCandidate:
    fragment_index: tuple[int, int]
    name: str
    value: str
    unit: str | None
    quote: str
    method: str
    confidence: float


def normalize_text(value: str) -> str:
    normalized = unicodedata.normalize("NFC", value)
    return " ".join(normalized.replace("\u00a0", " ").split())


def parse_document(data: bytes, media_type: str) -> ParsedDocument:
    try:
        if media_type == PDF_MEDIA_TYPE:
            return _parse_pdf(data)
        if media_type == DOCX_MEDIA_TYPE:
            return _parse_docx(data)
        if media_type == XLSX_MEDIA_TYPE:
            return _parse_xlsx(data)
    except NativeParseError:
        raise
    except Exception as exc:
        raise NativeParseError("DOCUMENT_PARSE_FAILED") from exc
    raise NativeParseError("UNSUPPORTED_DOCUMENT_FORMAT")


def extract_parameters(document: ParsedDocument) -> list[ParameterCandidate]:
    candidates: list[ParameterCandidate] = []
    seen: set[tuple[str, str, str | None, tuple[int, int]]] = set()
    for page_index, page in enumerate(document.pages):
        for fragment_index, fragment in enumerate(page.fragments):
            if page.active_layer == "OCR" and fragment.method != "OCR":
                continue
            if fragment.locator.get("pageWords"):
                continue
            candidate = _table_candidate(fragment, page_index, fragment_index)
            if candidate is None:
                candidate = _text_candidate(fragment, page_index, fragment_index)
            if candidate is None:
                continue
            key = (
                candidate.name.casefold(),
                candidate.value.casefold(),
                candidate.unit.casefold() if candidate.unit else None,
                candidate.fragment_index,
            )
            if key not in seen:
                seen.add(key)
                candidates.append(candidate)
    return candidates


def _parse_pdf(data: bytes) -> ParsedDocument:
    rich = _parse_pdf_with_pymupdf(data)
    if rich is not None:
        return rich
    reader = PdfReader(BytesIO(data), strict=False)
    if len(reader.pages) > MAX_PAGES:
        raise NativeParseError("DOCUMENT_PAGE_LIMIT_EXCEEDED")
    pages: list[ParsedPage] = []
    for page_no, source_page in enumerate(reader.pages, start=1):
        text = source_page.extract_text() or ""
        fragments = tuple(
            ParsedFragment(
                text=line,
                method="TEXT",
                locator_kind="PDF",
                locator={"page": page_no, "line": line_no},
            )
            for line_no, line in enumerate(_safe_lines(text), start=1)
        )
        pages.append(_page(page_no, fragments))
    return ParsedDocument(tuple(pages))


def _parse_pdf_with_pymupdf(data: bytes) -> ParsedDocument | None:
    try:
        import fitz  # type: ignore[import-not-found]
    except ImportError:
        return None
    try:
        document = fitz.open(stream=data, filetype="pdf")
    except Exception as exc:
        raise NativeParseError("PDF_PARSE_FAILED") from exc
    try:
        if document.page_count > MAX_PAGES:
            raise NativeParseError("DOCUMENT_PAGE_LIMIT_EXCEEDED")
        pages: list[ParsedPage] = []
        all_page_words: list[tuple[tuple, ...]] = []
        for page_index in range(document.page_count):
            source_page = document.load_page(page_index)
            page_no = page_index + 1
            words = _page_words(source_page)
            all_page_words.append(tuple(words))
            fragments: list[ParsedFragment] = []
            for line_no, line in enumerate(_lines_from_words(words), start=1):
                fragments.append(
                    ParsedFragment(
                        text=_bounded(line),
                        method="TEXT",
                        locator_kind="PDF",
                        locator={"page": page_no, "line": line_no},
                    )
                )

            page_context = _page_context(words)
            if page_context:
                fragments.append(
                    ParsedFragment(
                        text=_bounded(page_context),
                        method="TEXT",
                        locator_kind="PDF",
                        locator={"page": page_no, "pageWords": True},
                    )
                )

            rects = _page_rects(source_page)
            word_tables = detect_tables_on_page(words, page_index)
            grid_tables = detect_tables_from_grid(words, rects, page_index)
            for table in merge_page_tables(page_index, grid_tables, word_tables):
                header_context = _table_header_context(table.rows)
                for row_index, row in enumerate(table.rows):
                    cells = tuple(normalize_text(cell.text) for cell in row)
                    if not any(cells):
                        continue
                    row_text = _table_row_text(header_context, cells)
                    locator = {
                        "page": page_no,
                        "table": table.id,
                        "row": row_index + 1,
                    }
                    if table.bbox:
                        locator["tableBbox"] = table.bbox
                    if table.source:
                        locator["tableSource"] = table.source
                    fragments.append(
                        ParsedFragment(
                            text=_bounded(row_text),
                            method="TABLE",
                            locator_kind="PDF",
                            locator=locator,
                            cells=cells,
                            cell_boxes=tuple(cell.to_dict() for cell in row),
                            table_bbox=tuple(table.bbox or ()),
                            table_source=table.source,
                        )
                    )

            rect = source_page.rect
            pages.append(
                ParsedPage(
                    page_no=page_no,
                    fragments=tuple(fragments),
                    quality_score=_quality(fragments),
                    ocr_needed=sum(len(fragment.text) for fragment in fragments) < 20,
                    width_px=int(rect.width),
                    height_px=int(rect.height),
                )
            )
        return ParsedDocument(tuple(pages), tuple(all_page_words))
    finally:
        document.close()


def _page_words(page) -> list[tuple]:
    words = []
    for item in page.get_text("words") or []:
        if len(item) < 5:
            continue
        text = normalize_text(str(item[4]))
        if not text:
            continue
        words.append((float(item[0]), float(item[1]), float(item[2]), float(item[3]), text))
    return words


def _page_rects(page) -> list[tuple]:
    rects: list[tuple] = []
    for drawing in page.get_drawings() or []:
        for item in drawing.get("items", []):
            if not item:
                continue
            op = item[0]
            if op == "re" and len(item) >= 2:
                rect = item[1]
                rects.append((float(rect.x0), float(rect.y0), float(rect.x1), float(rect.y1)))
            elif op == "l" and len(item) >= 3:
                p1, p2 = item[1], item[2]
                rects.append((float(p1.x), float(p1.y), float(p2.x), float(p2.y)))
    return rects


def _lines_from_words(words: list[tuple]) -> list[str]:
    lines = []
    for line in group_words_into_lines(words):
        text = normalize_text(" ".join(word[4] for word in line))
        if text:
            lines.append(text)
    return lines


def _page_context(words: list[tuple]) -> str:
    return "\n".join(_lines_from_words(words))


def _table_header_context(rows) -> tuple[str, ...]:
    context: list[str] = []
    for row in rows[:4]:
        text = normalize_text(" | ".join(cell.text for cell in row if cell.text.strip()))
        if text:
            context.append(text)
    return tuple(context)


def _table_row_text(header_context: tuple[str, ...], cells: tuple[str, ...]) -> str:
    row = " | ".join(cells)
    if not header_context:
        return row
    return " || ".join((*header_context, row))


def _quality(fragments: list[ParsedFragment]) -> float:
    chars = sum(len(fragment.text) for fragment in fragments)
    return min(1.0, chars / 80) if chars else 0.0


def _parse_docx(data: bytes) -> ParsedDocument:
    document = DocxDocument(BytesIO(data))
    fragments: list[ParsedFragment] = []
    for paragraph_no, paragraph in enumerate(document.paragraphs, start=1):
        text = normalize_text(paragraph.text)
        if text:
            fragments.append(
                ParsedFragment(
                    text=_bounded(text),
                    method="TEXT",
                    locator_kind="DOCX",
                    locator={"paragraph": paragraph_no},
                )
            )
    for table_no, table in enumerate(document.tables, start=1):
        for row_no, row in enumerate(table.rows, start=1):
            cells = tuple(normalize_text(cell.text) for cell in row.cells)
            if not any(cells):
                continue
            text = _bounded(" | ".join(cells))
            fragments.append(
                ParsedFragment(
                    text=text,
                    method="TABLE",
                    locator_kind="DOCX",
                    locator={"table": table_no, "row": row_no},
                    cells=cells,
                )
            )
    return ParsedDocument((_page(1, tuple(fragments)),))


def _parse_xlsx(data: bytes) -> ParsedDocument:
    workbook = load_workbook(
        BytesIO(data), read_only=True, data_only=False, keep_links=False
    )
    pages: list[ParsedPage] = []
    try:
        for sheet_no, sheet in enumerate(workbook.worksheets, start=1):
            fragments: list[ParsedFragment] = []
            for row_no, row in enumerate(sheet.iter_rows(values_only=False), start=1):
                if row_no > MAX_ROWS_PER_SHEET:
                    raise NativeParseError("SPREADSHEET_ROW_LIMIT_EXCEEDED")
                values = tuple(_cell_text(cell.value) for cell in row)
                while values and not values[-1]:
                    values = values[:-1]
                if not any(values):
                    continue
                end_column = get_column_letter(max(1, len(values)))
                fragments.append(
                    ParsedFragment(
                        text=_bounded(" | ".join(values)),
                        method="TABLE",
                        locator_kind="XLSX",
                        locator={
                            "sheet": sheet.title,
                            "range": f"A{row_no}:{end_column}{row_no}",
                        },
                        cells=values,
                    )
                )
            pages.append(_page(sheet_no, tuple(fragments)))
    finally:
        workbook.close()
    return ParsedDocument(tuple(pages))


def _safe_lines(text: str) -> list[str]:
    lines = []
    for raw_line in text.splitlines():
        line = normalize_text(raw_line)
        if line:
            lines.append(_bounded(line))
    return lines


def _bounded(value: str) -> str:
    if len(value) > MAX_FRAGMENT_CHARS:
        raise NativeParseError("DOCUMENT_FRAGMENT_LIMIT_EXCEEDED")
    return value


def _page(page_no: int, fragments: tuple[ParsedFragment, ...]) -> ParsedPage:
    chars = sum(len(fragment.text) for fragment in fragments)
    quality = min(1.0, chars / 80) if chars else 0.0
    return ParsedPage(page_no, fragments, quality, chars < 20)


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    return normalize_text(str(value))


def _table_candidate(
    fragment: ParsedFragment, page_index: int, fragment_index: int
) -> ParameterCandidate | None:
    if fragment.cells is None:
        return None
    cells = tuple(cell for cell in fragment.cells if cell)
    if len(cells) < 2 or _looks_like_header(cells):
        return None
    name = cells[0].strip(" :-")
    if len(name) < 2:
        return None
    value, unit = _split_value_unit(cells[1], cells[2] if len(cells) > 2 else None)
    if not value:
        return None
    return ParameterCandidate(
        fragment_index=(page_index, fragment_index),
        name=name,
        value=value,
        unit=unit,
        quote=fragment.text,
        method="TABLE",
        confidence=(
            min(0.95, fragment.confidence)
            if fragment.confidence is not None
            else 0.95
        ),
    )


def _text_candidate(
    fragment: ParsedFragment, page_index: int, fragment_index: int
) -> ParameterCandidate | None:
    match = _PAIR.match(fragment.text)
    if match is None:
        return None
    name = match.group(1).strip(" .:-")
    if name.casefold() in _HEADER_WORDS:
        return None
    value, unit = _split_value_unit(match.group(2), None)
    if not name or not value:
        return None
    return ParameterCandidate(
        fragment_index=(page_index, fragment_index),
        name=name,
        value=value,
        unit=unit,
        quote=fragment.text,
        method="REGEX",
        confidence=(
            min(0.9, fragment.confidence)
            if fragment.confidence is not None
            else 0.9
        ),
    )


def _split_value_unit(value: str, explicit_unit: str | None) -> tuple[str, str | None]:
    cleaned = normalize_text(value).strip(" ;")
    unit = normalize_text(explicit_unit).strip(" ;") if explicit_unit else None
    if unit:
        return cleaned, unit
    match = _NUMBER_WITH_UNIT.match(cleaned)
    if match is None:
        return cleaned, None
    prefix = (match.group(1) or "").strip()
    numeric = match.group(2).strip()
    parsed_value = f"{prefix}{numeric}" if prefix else numeric
    parsed_unit = (match.group(3) or "").strip(" ;") or None
    return parsed_value, parsed_unit


def _looks_like_header(cells: tuple[str, ...]) -> bool:
    first = cells[0].casefold().strip()
    second = cells[1].casefold().strip()
    return first in _HEADER_WORDS or (
        any(word in first for word in ("параметр", "характерист", "наименован"))
        and any(word in second for word in ("значен", "value"))
    )
