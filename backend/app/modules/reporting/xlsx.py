"""Рендер XLSX из snapshot отчёта.

Читает только `payload` snapshot — ни базы, ни ORM: отчёт, выданный заказчику,
обязан собираться из зафиксированных данных, а не из текущего состояния
проверки.

Две вещи, ради которых этот файл существует отдельно:

* **Формулы.** Значение из документа поставщика попадает в ячейку как есть.
  Строка, начинающаяся с `=`, `+`, `-` или `@`, в Excel становится формулой —
  классический spreadsheet formula injection (reports §QA). Такие значения
  экранируются апострофом.
* **Статус словами и цветом.** Контракт требует и текст, и цвет: чёрно-белая
  печать и дальтонизм не должны лишать отчёт смысла.
"""

from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

#: Человеческие подписи статусов (канон §15.4). Отчёт читает менеджер и
#: заказчик, машинные коды им ни о чём не говорят.
STATUS_LABELS = {
    "COMPLIANT": "СООТВЕТСТВУЕТ",
    "NON_COMPLIANT": "НЕ СООТВЕТСТВУЕТ",
    "REVIEW_REQUIRED": "ТРЕБУЕТ ПРОВЕРКИ",
    "NOT_FOUND_IN_FACT": "НЕ НАЙДЕНО В ПАСПОРТЕ",
    "NOT_REQUIRED": "НЕ ТРЕБУЕТСЯ В ТЗ",
}

VERDICT_LABELS = {
    "COMPLIANT": "СООТВЕТСТВУЕТ",
    "PARTIAL": "ЧАСТИЧНО",
    "NON_COMPLIANT": "НЕ СООТВЕТСТВУЕТ",
    "PENDING": "НЕ ПОДТВЕРЖДЁН",
}

CRITICALITY_LABELS = {
    "BLOCKER": "КРИТИЧЕСКИЙ",
    "MAJOR": "ЗНАЧИМЫЙ",
    "MINOR": "СПРАВОЧНЫЙ",
}

_FILLS = {
    "COMPLIANT": "FFD8F3DC",
    "NON_COMPLIANT": "FFFAD5D5",
    "REVIEW_REQUIRED": "FFFFF0C2",
    "NOT_FOUND_IN_FACT": "FFFFE8CC",
    "NOT_REQUIRED": "FFEDEDED",
}

_COLUMNS = (
    ("Параметр", 34),
    ("Требование ТЗ", 20),
    ("Ед.", 10),
    ("Факт по паспорту", 20),
    ("Ед. факта", 10),
    ("Правило", 26),
    ("Статус", 24),
    ("Значимость", 14),
    ("Допуск", 18),
    ("Комментарий", 30),
)

#: Символы, с которых Excel начинает разбирать содержимое ячейки как формулу.
_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")


def render_batch(payload: dict) -> bytes:
    """Сводный отчёт заявки: лист «Сводка» и по листу на позицию.

    Отдельная функция, а не флаг в `render`: у сводного отчёта другой payload —
    заявка со списком позиций вместо одной позиции, — и смешивать их в одной
    ветке значило бы разбирать структуру прямо в рендере.
    """
    book = Workbook()
    summary = book.active
    summary.title = "Сводка"
    _render_batch_summary(summary, payload)

    for section in payload.get("items", []):
        sheet = book.create_sheet(_position_title(section))
        _render_item(sheet, section)
    return _save(book)


def _render_batch_summary(sheet, payload: dict) -> None:
    batch = payload.get("batch", {})
    verdict = payload.get("verdict", {})
    sections = payload.get("items", [])

    rows = [
        ("Заявка", batch.get("title")),
        ("Заказчик", batch.get("customer")),
        ("Поставщик", batch.get("supplier")),
        ("Позиций", batch.get("itemsTotal")),
        ("Вердикт системы", VERDICT_LABELS.get(verdict.get("computed"), "—")),
        ("Вердикт менеджера", VERDICT_LABELS.get(verdict.get("manager"), "—")),
        ("Комментарий менеджера", verdict.get("comment")),
        ("Подтверждён", verdict.get("approvedAt")),
        ("Отчёт сформирован", payload.get("generatedAt")),
    ]
    sheet["A1"] = "Сводный отчёт о проверке технических характеристик"
    sheet["A1"].font = Font(bold=True, size=14)
    for offset, (label, value) in enumerate(rows, start=3):
        sheet.cell(row=offset, column=1, value=label).font = Font(bold=True)
        sheet.cell(row=offset, column=2, value=_safe(value))

    start = 3 + len(rows) + 1
    sheet.cell(row=start, column=1, value="Позиции заявки").font = Font(bold=True)
    headers = ("№", "Позиция", "Поставщик", "Вердикт", *STATUS_LABELS.values())
    for index, title in enumerate(headers, start=1):
        cell = sheet.cell(row=start + 1, column=index, value=title)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(vertical="center", wrap_text=True)

    for offset, section in enumerate(sections, start=start + 2):
        counts = _status_counts(section)
        item_verdict = (section.get("verdict") or {}).get("manager")
        if item_verdict in (None, "PENDING"):
            item_verdict = (section.get("verdict") or {}).get("computed")
        values = (
            section.get("positionNo"),
            section.get("name"),
            section.get("supplier"),
            VERDICT_LABELS.get(item_verdict, "—"),
            *(counts[status] for status in STATUS_LABELS),
        )
        for index, value in enumerate(values, start=1):
            sheet.cell(row=offset, column=index, value=_safe(value))

    for index, width in enumerate((6, 34, 24, 22, 16, 18, 18, 20, 20), start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width


def _status_counts(section: dict) -> dict[str, int]:
    counts = dict.fromkeys(STATUS_LABELS, 0)
    for row in section.get("rows", []):
        status = row.get("status")
        if status in counts:
            counts[status] += 1
    return counts


def _position_title(section: dict) -> str:
    """Имя листа позиции: номер плюс название, в пределах 31 символа Excel."""
    number = section.get("positionNo")
    name = str(section.get("name") or "").strip()
    raw = f"{number}. {name}" if number else name or "Позиция"
    for symbol in "[]:*?/\\":
        raw = raw.replace(symbol, " ")
    return raw.strip()[:31] or "Позиция"


def render(payload: dict) -> bytes:
    """Собирает книгу: лист «Сводка» и лист позиции."""
    book = Workbook()
    summary = book.active
    summary.title = "Сводка"
    _render_summary(summary, payload)
    _render_item(book.create_sheet(_sheet_title(payload)), payload)
    return _save(book)


def _save(book: Workbook) -> bytes:
    target = BytesIO()
    book.save(target)
    return target.getvalue()


def _render_summary(sheet, payload: dict) -> None:
    item = payload.get("item", {})
    verdict = payload.get("verdict", {})
    versions = payload.get("versions", {})
    basis = verdict.get("basis") or {}

    rows = [
        ("Проверка", item.get("batchTitle")),
        ("Заказчик", item.get("customer")),
        ("Позиция", item.get("name")),
        ("Поставщик", item.get("supplier")),
        ("Вердикт системы", VERDICT_LABELS.get(verdict.get("computed"), "—")),
        ("Вердикт менеджера", VERDICT_LABELS.get(verdict.get("manager"), "—")),
        ("Комментарий менеджера", verdict.get("comment")),
        ("Подтверждён", verdict.get("approvedAt")),
        ("Отчёт сформирован", payload.get("generatedAt")),
        ("Версия каталога", versions.get("catalog")),
        ("Версия правил", versions.get("rules")),
        ("Версия pipeline", versions.get("pipeline")),
    ]
    sheet["A1"] = "Отчёт о проверке технических характеристик"
    sheet["A1"].font = Font(bold=True, size=14)
    for offset, (label, value) in enumerate(rows, start=3):
        sheet.cell(row=offset, column=1, value=label).font = Font(bold=True)
        sheet.cell(row=offset, column=2, value=_safe(value))

    start = 3 + len(rows) + 1
    sheet.cell(row=start, column=1, value="Распределение статусов").font = Font(
        bold=True
    )
    for offset, (status, label) in enumerate(STATUS_LABELS.items(), start=start + 1):
        sheet.cell(row=offset, column=1, value=label)
        sheet.cell(row=offset, column=2, value=int(basis.get(status, 0)))

    sheet.column_dimensions["A"].width = 30
    sheet.column_dimensions["B"].width = 46


def _render_item(sheet, payload: dict) -> None:
    for index, (title, width) in enumerate(_COLUMNS, start=1):
        cell = sheet.cell(row=1, column=index, value=title)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        sheet.column_dimensions[get_column_letter(index)].width = width

    for offset, row in enumerate(payload.get("rows", []), start=2):
        status = row.get("status")
        values = (
            row.get("parameter"),
            row.get("requirementValue"),
            row.get("requirementUnit"),
            row.get("factValue"),
            row.get("factUnit"),
            row.get("expression"),
            STATUS_LABELS.get(status, status),
            CRITICALITY_LABELS.get(row.get("criticality"), row.get("criticality")),
            row.get("tolerance"),
            row.get("comment"),
        )
        for index, value in enumerate(values, start=1):
            cell = sheet.cell(row=offset, column=index, value=_safe(value))
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        fill = _FILLS.get(status)
        if fill:
            sheet.cell(row=offset, column=7).fill = PatternFill(
                start_color=fill, end_color=fill, fill_type="solid"
            )
    sheet.freeze_panes = "A2"


def _sheet_title(payload: dict) -> str:
    """Имя листа позиции: Excel запрещает часть символов и длину свыше 31."""
    raw = str(payload.get("item", {}).get("name") or "Позиция")
    for symbol in "[]:*?/\\":
        raw = raw.replace(symbol, " ")
    return raw.strip()[:31] or "Позиция"


def _safe(value) -> str | int | float | None:
    """Экранирует значение, которое Excel принял бы за формулу."""
    if value is None or isinstance(value, (int, float)):
        return value
    text = str(value)
    if text.startswith(_FORMULA_PREFIXES):
        return "'" + text
    return text
