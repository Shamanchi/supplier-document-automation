"""Workbook loading: xlsx/xls -> per-sheet 2D grid of cells.

Merged cells are filled: every covered cell gets the master value, but
only the anchor keeps is_merged_fill=False. Row classification relies on
that flag so a category caption merged across the whole row is not
mistaken for a data row with prices.
"""
from __future__ import annotations

import io
from dataclasses import dataclass, field
from decimal import Decimal, ROUND_HALF_UP
import re
import zipfile

MAX_ROWS = 20000
MAX_COLS = 64


@dataclass
class Cell:
    value: object | None = None
    # Preserve the source number format because it may carry otherwise absent
    # monetary evidence (for example ``#,##0\"р.\"`` for roubles).
    number_format: str | None = None
    # Spreadsheet applications may display a rounded integer while retaining a
    # fractional number underneath.  Keep that visible value separately so a
    # price parser can honour what the supplier actually published.
    display_value: str | None = None
    is_merged_fill: bool = False
    # Formatting is semantic evidence in supplier price lists: bold rows are
    # commonly category captions separating several price fragments that all
    # reuse one column header. It is never used to assign a column role by
    # itself, only to find safe section boundaries.
    is_bold: bool = False
    # Filled, bold, merged rows are common section captions in XLSX price
    # lists.  Preserve the fill kind so deterministic multi-panel layouts can
    # distinguish those captions from ordinary notes without depending on
    # supplier text or coordinates.
    fill_type: str | None = None
    fill_tint: float | None = None


@dataclass
class SheetGrid:
    name: str
    rows: list[list[Cell]] = field(default_factory=list)

    @property
    def n_cols(self) -> int:
        return max((len(r) for r in self.rows), default=0)


def is_empty_grid(grid: SheetGrid) -> bool:
    """Return whether a worksheet has no meaningful cell values.

    Workbooks commonly contain empty placeholder tabs.  They are neither a
    price table nor a parsing problem, so callers must not report them as a
    rejected sheet.
    """
    for row in grid.rows:
        for cell in row:
            value = cell.value
            if value is None:
                continue
            if isinstance(value, str) and not value.strip():
                continue
            return False
    return True


_QUOTED_FORMAT_TEXT_RE = re.compile(r'"[^\"]*"')
_BRACKETED_FORMAT_PART_RE = re.compile(r"\[[^\]]*\]")


def _fraction_places(number_format: str | None) -> int | None:
    """Return the visible decimal precision of a simple Excel number format.

    ``None`` means that the format is General or does not look numeric.  We
    only need the zero-precision case here: it tells the parser that a supplier
    intentionally published rounded roubles rather than hidden kopecks.
    """
    if not number_format or number_format.lower() == "general":
        return None
    section = number_format.split(";", 1)[0]
    section = _QUOTED_FORMAT_TEXT_RE.sub("", section)
    section = _BRACKETED_FORMAT_PART_RE.sub("", section)
    if not any(char in "0#" for char in section):
        return None
    dot = re.search(r"\.(?P<fraction>[0#]+)", section)
    if dot:
        return len(dot.group("fraction"))
    # In locale-specific formats a comma can be the decimal separator.  It is
    # only treated as one when it follows a required digit, so ``#,##0`` keeps
    # its conventional thousands separator meaning.
    comma = re.search(r"(?<=0),(?P<fraction>[0#]+)", section)
    if comma:
        return len(comma.group("fraction"))
    return 0


def _display_numeric_value(value: object, number_format: str | None) -> str | None:
    if isinstance(value, bool) or not isinstance(value, (int, float, Decimal)):
        return None
    if _fraction_places(number_format) != 0:
        return None
    rounded = Decimal(str(value)).quantize(Decimal("1"), rounding=ROUND_HALF_UP)
    return str(int(rounded))


def load_grids(data: bytes, filename: str) -> list[SheetGrid]:
    lower = filename.lower()
    if lower.endswith(".xls"):
        return _load_xls(data)
    return _load_xlsx(data)


def _load_xlsx(data: bytes) -> list[SheetGrid]:
    import openpyxl

    data = _normalize_ooxml_shared_strings_path(data)
    wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=False)
    formula_wb = openpyxl.load_workbook(
        io.BytesIO(data), data_only=False, read_only=False
    )
    grids: list[SheetGrid] = []
    for ws in wb.worksheets:
        if ws.sheet_state != "visible":
            continue
        formula_ws = formula_wb[ws.title]
        max_row = min(ws.max_row or 0, MAX_ROWS)
        max_col = min(ws.max_column or 0, MAX_COLS)
        rows: list[list[Cell]] = []
        for row in ws.iter_rows(min_row=1, max_row=max_row, max_col=max_col):
            loaded_row: list[Cell] = []
            for source_cell in row:
                formula_cell = formula_ws[source_cell.coordinate]
                semantic_empty = _cached_formula_zero_from_blank_inputs(
                    cached_value=source_cell.value,
                    formula=formula_cell.value,
                    formula_ws=formula_ws,
                )
                value = None if semantic_empty else source_cell.value
                loaded_row.append(
                    Cell(
                        value=value,
                        number_format=source_cell.number_format,
                        display_value=_display_numeric_value(
                            value, source_cell.number_format
                        ),
                        is_bold=bool(source_cell.font and source_cell.font.bold),
                        fill_type=(
                            source_cell.fill.fill_type
                            if source_cell.fill is not None
                            else None
                        ),
                        fill_tint=(
                            float(source_cell.fill.fgColor.tint)
                            if source_cell.fill is not None
                            and source_cell.fill.fgColor is not None
                            else None
                        ),
                    )
                )
            rows.append(loaded_row)
        # Fill merged ranges with the anchor value.
        for rng in ws.merged_cells.ranges:
            r0, c0 = rng.min_row - 1, rng.min_col - 1
            if r0 >= len(rows) or c0 >= len(rows[r0]):
                continue
            master_cell = rows[r0][c0]
            for r in range(r0, min(rng.max_row, max_row)):
                for c in range(c0, min(rng.max_col, max_col)):
                    if r == r0 and c == c0:
                        continue
                    if r < len(rows) and c < len(rows[r]):
                        rows[r][c] = Cell(
                            value=master_cell.value,
                            number_format=master_cell.number_format,
                            display_value=master_cell.display_value,
                            is_merged_fill=True,
                            is_bold=master_cell.is_bold,
                            fill_type=master_cell.fill_type,
                            fill_tint=master_cell.fill_tint,
                        )
        grids.append(SheetGrid(name=ws.title, rows=rows))
    wb.close()
    formula_wb.close()
    return grids


def _cached_formula_zero_from_blank_inputs(
    *,
    cached_value: object | None,
    formula: object | None,
    formula_ws,
) -> bool:
    """Treat a derived cached zero as empty when every formula input is empty.

    Excel stores cached formula results in the workbook.  A formula such as
    ``=C6+C6*35%`` therefore looks like a literal zero to ``data_only=True``
    after the base price in C6 has been cleared.  It is not an offered zero
    price.  We only suppress the narrow, provable case: a numeric cached zero,
    a same-sheet formula with at least one cell reference, and every referenced
    input blank.  Literal zeroes and formulae with real inputs are retained.
    """
    if isinstance(cached_value, bool) or cached_value not in (0, 0.0, Decimal("0")):
        return False
    if not isinstance(formula, str) or not formula.startswith("="):
        return False

    from openpyxl.formula import Tokenizer
    from openpyxl.utils.cell import range_boundaries

    references: list[str] = []
    try:
        for token in Tokenizer(formula).items:
            if token.type != "OPERAND" or token.subtype != "RANGE":
                continue
            reference = token.value
            # Cross-sheet, named-range and external references are deliberately
            # left untouched: their emptiness cannot be proved locally.
            if "!" in reference or "[" in reference or "]" in reference:
                return False
            reference = reference.replace("$", "")
            try:
                min_col, min_row, max_col, max_row = range_boundaries(reference)
            except ValueError:
                return False
            for row_index in range(min_row, max_row + 1):
                for column_index in range(min_col, max_col + 1):
                    references.append(formula_ws.cell(row_index, column_index).coordinate)
    except Exception:
        return False
    if not references:
        return False
    return all(
        formula_ws[coordinate].value is None
        or (
            isinstance(formula_ws[coordinate].value, str)
            and not formula_ws[coordinate].value.strip()
        )
        for coordinate in references
    )


def _normalize_ooxml_shared_strings_path(data: bytes) -> bytes:
    """Add the canonical shared-strings member for case-broken OOXML exports.

    OPC part names are case-sensitive, but several otherwise readable supplier
    exports store ``xl/SharedStrings.xml`` while their relationship targets
    ``xl/sharedStrings.xml``.  Excel accepts those files on a case-insensitive
    filesystem; openpyxl on Linux correctly does not.  Copy only the known
    OOXML part to its canonical spelling so the workbook remains unchanged in
    every other respect.
    """
    canonical = "xl/sharedStrings.xml"
    with zipfile.ZipFile(io.BytesIO(data)) as source:
        names = source.namelist()
        if canonical in names:
            return data
        match = next((name for name in names if name.casefold() == canonical.casefold()), None)
        if match is None:
            return data

        # ``ZipFile.writestr`` mutates the supplied ZipInfo's header offset,
        # so read every source member before reusing its metadata.
        member_payloads = [(member, source.read(member)) for member in source.infolist()]
        shared_strings = source.read(match)
        normalized = io.BytesIO()
        with zipfile.ZipFile(normalized, mode="w") as destination:
            for member, payload in member_payloads:
                destination.writestr(member, payload)
            destination.writestr(canonical, shared_strings)
        return normalized.getvalue()


def _load_xls(data: bytes) -> list[SheetGrid]:
    import xlrd

    wb = xlrd.open_workbook(file_contents=data, formatting_info=True)
    grids: list[SheetGrid] = []
    for ws in wb.sheets():
        max_row = min(ws.nrows, MAX_ROWS)
        max_col = min(ws.ncols, MAX_COLS)
        rows: list[list[Cell]] = []
        for r in range(max_row):
            row: list[Cell] = []
            for c in range(max_col):
                v = ws.cell_value(r, c)
                if v == "":
                    v = None
                number_format = None
                if v is not None and ws.cell_type(r, c) == xlrd.XL_CELL_NUMBER:
                    xf = wb.xf_list[ws.cell_xf_index(r, c)]
                    number_format = wb.format_map[xf.format_key].format_str
                row.append(
                    Cell(
                        value=v,
                        number_format=number_format,
                        display_value=_display_numeric_value(v, number_format),
                        is_bold=bool(
                            wb.font_list[
                                wb.xf_list[ws.cell_xf_index(r, c)].font_index
                            ].bold
                        ),
                    )
                )
            rows.append(row)
        for r0, r1, c0, c1 in ws.merged_cells:
            if r0 >= max_row or c0 >= max_col:
                continue
            master_cell = rows[r0][c0]
            for r in range(r0, min(r1, max_row)):
                for c in range(c0, min(c1, max_col)):
                    if r == r0 and c == c0:
                        continue
                    rows[r][c] = Cell(
                        value=master_cell.value,
                        number_format=master_cell.number_format,
                        display_value=master_cell.display_value,
                        is_merged_fill=True,
                        is_bold=master_cell.is_bold,
                        fill_type=master_cell.fill_type,
                        fill_tint=master_cell.fill_tint,
                    )
        grids.append(SheetGrid(name=ws.name, rows=rows))
    return grids
