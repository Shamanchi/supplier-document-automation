"""Human-auditable, persistence-free exports of a rule dry-run."""
from __future__ import annotations

from collections import OrderedDict
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.parsing.pipeline import ParsedItem


def _header_cells(item: ParsedItem) -> list[dict[str, Any]]:
    return [
        {"column_index": int(cell[0]), "value": str(cell[1])}
        for cell in (item.source_header_cells or [])
        if isinstance(cell, list) and len(cell) >= 2
    ]


def _price_payload(value) -> dict[str, Any]:
    return {
        "kind": value.kind,
        "amount": str(value.amount) if value.amount is not None else None,
        "amount_from": str(value.amount_from) if value.amount_from is not None else None,
        "amount_to": str(value.amount_to) if value.amount_to is not None else None,
        "qualifier": value.qualifier,
        "is_on_request": value.is_on_request,
        "vat_mode": value.vat_mode,
        "currency": value.currency,
        "min_quantity": (
            str(value.min_quantity) if value.min_quantity is not None else None
        ),
        "max_quantity": (
            str(value.max_quantity) if value.max_quantity is not None else None
        ),
        "basis_quantity": (
            str(value.basis_quantity) if value.basis_quantity is not None else None
        ),
        "basis_unit": value.basis_unit,
        "source_label": value.source_label,
        "source_column_index": value.source_column_index,
        "raw_value": value.raw_value,
    }


def _item_payload(item: ParsedItem) -> dict[str, Any]:
    return {
        "source_sheet": item.source_sheet,
        "source_row": item.source_row,
        "source_section_id": item.source_section_id,
        "source_panel_index": item.source_panel_index,
        "name": item.source_name,
        "normalized_name": item.normalized_name,
        "article": item.source_article,
        "unit": item.unit_code,
        "category": item.category_path,
        "dimensions": item.dimensions,
        "description": item.description,
        "attributes": [
            {
                "code": attribute.attribute_code,
                "name": attribute.source_header,
                "raw_value": attribute.raw_value,
                "value_kind": attribute.value_kind,
                "unit": attribute.canonical_unit,
                "source_column_index": attribute.source_column_index,
            }
            for attribute in item.attribute_values
        ],
        "prices": [_price_payload(value) for value in item.values],
        "source_header": _header_cells(item),
        "source_data": item.source_data_cells or [],
        "source_categories": item.source_category_cells or [],
    }


def build_validation_export(
    items: list[ParsedItem], *, filename: str
) -> dict[str, Any]:
    grouped: OrderedDict[tuple[str, str | None, str | None], list[ParsedItem]] = OrderedDict()
    for item in sorted(
        items,
        key=lambda candidate: (
            candidate.source_sheet,
            candidate.source_row,
            candidate.source_panel_index or 0,
            candidate.source_name,
        ),
    ):
        grouped.setdefault(
            (item.source_sheet, item.category_path, item.source_section_id), []
        ).append(item)
    sections = []
    for (sheet, category, section_id), section_items in grouped.items():
        source_header = next(
            (header for item in section_items if (header := _header_cells(item))), []
        )
        sections.append(
            {
                "source_sheet": sheet,
                "category": category,
                "source_section_id": section_id,
                "source_header": source_header,
                "position_count": len(section_items),
                "positions": [_item_payload(item) for item in section_items],
            }
        )
    return {"filename": filename, "item_count": len(items), "sections": sections}


def _attributes_text(item: dict[str, Any]) -> str:
    return "; ".join(
        f"{attribute['name']}: {attribute['raw_value']}"
        + (f" {attribute['unit']}" if attribute.get("unit") else "")
        for attribute in item["attributes"]
    )


def _prices_text(item: dict[str, Any]) -> str:
    rows = []
    for price in item["prices"]:
        if price["is_on_request"]:
            amount = "по запросу"
        elif price.get("qualifier") == "range" or price["amount_to"] is not None:
            amount = f"{price['amount_from'] or ''}–{price['amount_to'] or ''}"
        elif price.get("qualifier") == "from":
            amount = f"от {price['amount_from'] or price['amount'] or ''}".rstrip()
        elif price["amount"] is not None:
            amount = price["amount"]
        elif price["amount_from"] is not None:
            amount = price["amount_from"]
        else:
            amount = price["raw_value"] or "—"
        currency = f" {price['currency']}" if price.get("currency") else ""
        basis = (
            f" за {price['basis_quantity']} {price['basis_unit']}"
            if price.get("basis_quantity") is not None and price.get("basis_unit")
            else ""
        )
        label = price.get("source_label") or price["kind"]
        rows.append(f"{label}: {amount}{currency}{basis}")
    return "; ".join(rows)


def build_validation_xlsx(items: list[ParsedItem], *, filename: str) -> bytes:
    payload = build_validation_export(items, filename=filename)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Результат разбора"
    headers = [
        "Лист", "Строка", "Категория", "Наименование", "Артикул",
        "Единица", "Характеристики", "Цены", "Описание", "Секция правила",
    ]
    section_fill = PatternFill("solid", fgColor="D9EAF7")
    header_fill = PatternFill("solid", fgColor="EDEDED")
    for section_index, section in enumerate(payload["sections"]):
        if section_index:
            sheet.append([None] * len(headers))
        sheet.append(["Категория", section["category"] or "—", "Лист", section["source_sheet"]])
        for cell in sheet[sheet.max_row]:
            cell.font = Font(bold=True)
            cell.fill = section_fill
        sheet.append(headers)
        for cell in sheet[sheet.max_row]:
            cell.font = Font(bold=True)
            cell.fill = header_fill
        for item in section["positions"]:
            sheet.append(
                [
                    item["source_sheet"], item["source_row"], item["category"],
                    item["name"], item["article"], item["unit"],
                    _attributes_text(item), _prices_text(item), item["description"],
                    item["source_section_id"],
                ]
            )
    sheet.freeze_panes = "A3"
    sheet.auto_filter.ref = f"A2:J{sheet.max_row}"
    for index, cells in enumerate(sheet.iter_cols(), start=1):
        content_width = max(
            (
                max((len(line) for line in str(cell.value).splitlines()), default=0)
                for cell in cells
                if cell.value is not None
            ),
            default=0,
        )
        sheet.column_dimensions[get_column_letter(index)].width = min(
            60, max(8, content_width + 2)
        )
    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
