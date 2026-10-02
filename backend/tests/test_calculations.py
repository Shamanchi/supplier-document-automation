"""Скидки и расчёт: выбор правила, объёмный порог, ручная правка с причиной."""

from __future__ import annotations

import asyncio
import io
import uuid
from datetime import date
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from app.matching.calculations import CalculationService, LineInput, split_vat
from app.matching.discounts import DiscountService, PositionKey, apply_percent
from app.matching.export import build_quote_workbook
from app.matching.physical import import_reference
from app.models import (
    CanonicalProduct,
    DocumentLink,
    DiscountRule,
    DiscountScope,
    GlobalCategory,
    PriceKind,
    ProductDocument,
    SupplierPriceValue,
)
from tests.test_matching import PROJECT_ID, add_item, make_version, open_session


async def add_price(
    db, item, amount: str, *, kind: str = PriceKind.dealer.value, min_quantity=None
) -> SupplierPriceValue:
    value = SupplierPriceValue(
        item_id=item.id,
        price_kind=kind,
        amount=Decimal(amount),
        currency="RUB",
        min_quantity=Decimal(min_quantity) if min_quantity is not None else None,
    )
    db.add(value)
    await db.flush()
    return value


def rule(supplier_id, scope: str, percent: str, **kwargs) -> DiscountRule:
    return DiscountRule(
        project_id=PROJECT_ID,
        supplier_id=supplier_id,
        scope=scope,
        percent=Decimal(percent),
        **kwargs,
    )


# ------------------------------------------------------------------ выбор правила


def test_specific_rule_beats_general_even_when_smaller():
    """Артикульная скидка перекрывает поставщицкую, даже если процент ниже.

    Иначе точечная договорённость по конкретной позиции молча растворялась бы
    в общем проценте поставщика.
    """
    supplier = uuid.uuid4()
    rules = [
        rule(supplier, DiscountScope.supplier.value, "20"),
        rule(supplier, DiscountScope.article.value, "5", normalized_article="ОП-4"),
    ]
    match = DiscountService(None).match(
        PositionKey(supplier_id=supplier, normalized_article="ОП-4"), rules
    )
    assert match.source == DiscountScope.article.value
    assert match.percent == Decimal("5")


def test_max_wins_inside_one_level():
    supplier = uuid.uuid4()
    category = uuid.uuid4()
    rules = [
        rule(supplier, DiscountScope.category.value, "7", category_id=category),
        rule(supplier, DiscountScope.category.value, "12", category_id=category),
    ]
    match = DiscountService(None).match(
        PositionKey(supplier_id=supplier, category_id=category), rules
    )
    assert match.percent == Decimal("12")


def test_rules_of_other_supplier_do_not_apply():
    mine, other = uuid.uuid4(), uuid.uuid4()
    rules = [rule(other, DiscountScope.supplier.value, "30")]
    match = DiscountService(None).match(PositionKey(supplier_id=mine), rules)
    assert match.percent == Decimal("0")
    assert match.source == "none"


def test_inactive_rule_is_ignored():
    supplier = uuid.uuid4()
    rules = [rule(supplier, DiscountScope.supplier.value, "10", is_active=False)]
    assert DiscountService(None).match(PositionKey(supplier_id=supplier), rules).percent == 0


def test_volume_discount_needs_the_threshold():
    """Supplier B: 4% от 300 тыс. ₽ — ниже порога скидки нет."""
    supplier = uuid.uuid4()
    rules = [
        rule(supplier, DiscountScope.volume.value, "4", min_amount=Decimal("300000")),
    ]
    service = DiscountService(None)
    assert service.volume(supplier, Decimal("299999"), rules).percent == Decimal("0")
    assert service.volume(supplier, Decimal("300000"), rules).percent == Decimal("4")


def test_apply_percent_rounds_to_kopecks():
    assert apply_percent(Decimal("1000.00"), Decimal("2.5")) == Decimal("975.00")
    assert apply_percent(Decimal("333.33"), Decimal("7")) == Decimal("310.00")
    assert apply_percent(None, Decimal("5")) is None


def test_split_vat_included_extracts_vat_out_of_price():
    """618.40 с НДС 22% (included): без НДС ~506.89, НДС ~111.51."""
    assert split_vat(Decimal("618.40"), Decimal("22"), "included") == (
        Decimal("506.89"),
        Decimal("111.51"),
        Decimal("618.40"),
    )


def test_split_vat_excluded_adds_vat_on_top():
    assert split_vat(Decimal("618.40"), Decimal("22"), "excluded") == (
        Decimal("618.40"),
        Decimal("136.05"),
        Decimal("754.45"),
    )


def test_split_vat_unknown_is_fail_closed():
    """P0-2: unknown не должен давать with_vat=754.45 — НДС сверху не начисляем."""
    assert split_vat(Decimal("618.40"), Decimal("22"), "unknown") == (
        Decimal("618.40"),
        None,
        Decimal("618.40"),
    )


def test_split_vat_none_mode_is_fail_closed_like_unknown():
    assert split_vat(Decimal("618.40"), Decimal("22"), None) == (
        Decimal("618.40"),
        None,
        Decimal("618.40"),
    )


def test_split_vat_without_rate_reports_no_vat_component():
    assert split_vat(Decimal("618.40"), None, "unknown") == (
        Decimal("618.40"),
        None,
        Decimal("618.40"),
    )


# ------------------------------------------------------------------ расчёт


def test_calculation_applies_supplier_discount(tmp_path):
    async def scenario():
        async with open_session(tmp_path) as db:
            version = await make_version(db, supplier_name="Supplier A")
            item = await add_item(db, version, "Огнетушитель ОП-4(з)", article="ОП-4")
            await add_price(db, item, "1000.00")
            db.add(rule(version.supplier_id, DiscountScope.supplier.value, "10"))
            await db.flush()

            service = CalculationService(db)
            calculation = await service.create(
                project_id=PROJECT_ID,
                lines=[LineInput(price_item_id=item.id, quantity=Decimal("3"))],
            )
            totals = await service.recalculate(calculation)
            row = calculation.items[0]
            assert row.base_price == Decimal("1000.00")
            assert row.discount_percent == Decimal("10")
            assert row.price_after_discount == Decimal("900.00")
            assert totals.base == Decimal("3000.00")
            assert totals.final == Decimal("2700.00")

    asyncio.run(scenario())


def test_quote_pipeline_uses_reference_and_exports_confirmed_certificate(tmp_path):
    """Сквозной путь T-21: позиция → физика → скидка → сертификат → Excel."""

    async def scenario():
        async with open_session(tmp_path) as db:
            version = await make_version(db, supplier_name="Supplier A")
            item = await add_item(db, version, "Огнетушитель ОП-4(з)", article="ОП-4")
            await add_price(db, item, "1000.00")
            await import_reference(
                db,
                project_id=PROJECT_ID,
                filename="ППО.xlsx.csv",
                payload="Артикул;Наименование;Вес, кг;Объём, м3\nОП-4;Огнетушитель;6,5;0,018\n".encode(),
            )
            canonical = CanonicalProduct(
                match_key=f"test-{uuid.uuid4()}",
                title=item.source_name,
                normalized_title=item.normalized_name,
                normalized_article=item.normalized_article,
            )
            document = ProductDocument(
                title="Сертификат ОП-4",
                registration_number="ЕАЭС RU С-RU.ПБ01.В.00001/26",
                valid_until=date(2028, 1, 1),
                storage_key="https://documents.example/cert-op4.pdf",
            )
            db.add_all([canonical, document])
            await db.flush()
            db.add(DocumentLink(
                document_id=document.id,
                canonical_product_id=canonical.id,
                confirmed_by="manager@example.test",
                evidence={"verdict": "valid", "valid_until": "2028-01-01"},
            ))
            await db.flush()

            calculation = await CalculationService(db).create(
                project_id=PROJECT_ID,
                lines=[LineInput(
                    price_item_id=item.id,
                    quantity=Decimal("2"),
                    canonical_id=canonical.id,
                    source_line="=2+2",
                )],
            )
            row = calculation.items[0]
            assert row.weight_kg == Decimal("6.5")
            assert row.volume_m3 == Decimal("0.018")
            assert row.weight_source == row.volume_source == "reference"
            assert row.certificate_confirmed_by == "manager@example.test"

            workbook = load_workbook(io.BytesIO(build_quote_workbook(calculation)))
            sheet = workbook.active
            headings = [cell.value for cell in sheet[1]]
            assert sheet.cell(2, headings.index("Вес итого, кг") + 1).value == 13
            assert sheet.cell(2, headings.index("Объём итого, м³") + 1).value == 0.036
            cert_cell = sheet.cell(2, headings.index("Ссылка/файл сертификата") + 1)
            assert cert_cell.hyperlink.target == "https://documents.example/cert-op4.pdf"
            source_cell = sheet.cell(2, headings.index("Исходная строка") + 1)
            assert source_cell.value == "'=2+2"
            assert source_cell.data_type == "s"
            workbook.close()

    asyncio.run(scenario())


def test_passport_without_verdict_fills_calculation_and_export(tmp_path):
    """Подтверждённый паспорт без вердикта FSA виден в позиции расчёта и в КП."""

    async def scenario():
        async with open_session(tmp_path) as db:
            version = await make_version(db, supplier_name="Supplier A")
            item = await add_item(db, version, "Огнетушитель ОП-5(з)", article="ОП-5")
            await add_price(db, item, "1000.00")
            canonical = CanonicalProduct(
                match_key=f"test-{uuid.uuid4()}",
                title=item.source_name,
                normalized_title=item.normalized_name,
                normalized_article=item.normalized_article,
            )
            passport = ProductDocument(
                kind="passport",
                title="Огнетушители порошковые переносные закачные ОП-1…ОП-5",
                storage_key="OCR_prices/Supplier A/Паспорта/ОП-1-5.pdf",
                source="passport_extraction",
            )
            db.add_all([canonical, passport])
            await db.flush()
            db.add(
                DocumentLink(
                    document_id=passport.id,
                    canonical_product_id=canonical.id,
                    matched_by="passport_extraction",
                    confirmed_by="demo_pipeline_wave2",
                    evidence={"source": "pdf_extraction"},
                )
            )
            await db.flush()

            calculation = await CalculationService(db).create(
                project_id=PROJECT_ID,
                lines=[
                    LineInput(
                        price_item_id=item.id,
                        quantity=Decimal("1"),
                        canonical_id=canonical.id,
                    )
                ],
            )
            row = calculation.items[0]
            assert row.certificate_number == passport.title
            assert row.certificate_reference == passport.storage_key
            assert row.certificate_confirmed_by == "demo_pipeline_wave2"

            workbook = load_workbook(io.BytesIO(build_quote_workbook(calculation)))
            sheet = workbook.active
            headings = [cell.value for cell in sheet[1]]
            cert_cell = sheet.cell(2, headings.index("Ссылка/файл сертификата") + 1)
            assert cert_cell.value == passport.storage_key
            assert cert_cell.hyperlink is None
            workbook.close()

    asyncio.run(scenario())


def test_valid_certificate_wins_over_passport_fallback(tmp_path):
    """Если у канона есть действующий сертификат и паспорт — в расчёт идёт сертификат."""

    async def scenario():
        async with open_session(tmp_path) as db:
            version = await make_version(db, supplier_name="Supplier A")
            item = await add_item(db, version, "Огнетушитель ОП-5(з)", article="ОП-5")
            await add_price(db, item, "1000.00")
            canonical = CanonicalProduct(
                match_key=f"test-{uuid.uuid4()}",
                title=item.source_name,
                normalized_title=item.normalized_name,
                normalized_article=item.normalized_article,
            )
            passport = ProductDocument(
                kind="passport",
                title="Паспорт ОП-5",
                storage_key="/files/passport-op5.pdf",
                source="passport_extraction",
            )
            certificate = ProductDocument(
                kind="certificate",
                title="Сертификат соответствия ОП-5",
                registration_number="ЕАЭС RU С-RU.ПБ01.В.00002/26",
                valid_until=date(2029, 1, 1),
                storage_key="https://documents.example/cert-op5.pdf",
            )
            db.add_all([canonical, passport, certificate])
            await db.flush()
            db.add(
                DocumentLink(
                    document_id=passport.id,
                    canonical_product_id=canonical.id,
                    confirmed_by="demo_pipeline_wave2",
                )
            )
            db.add(
                DocumentLink(
                    document_id=certificate.id,
                    canonical_product_id=canonical.id,
                    confirmed_by="manager@example.test",
                    evidence={"verdict": "valid"},
                )
            )
            await db.flush()

            calculation = await CalculationService(db).create(
                project_id=PROJECT_ID,
                lines=[
                    LineInput(
                        price_item_id=item.id,
                        quantity=Decimal("1"),
                        canonical_id=canonical.id,
                    )
                ],
            )
            row = calculation.items[0]
            assert row.certificate_number == certificate.registration_number
            assert row.certificate_valid_until == certificate.valid_until
            assert row.certificate_reference == certificate.storage_key
            assert row.certificate_confirmed_by == "manager@example.test"

    asyncio.run(scenario())


def test_expired_certificate_does_not_hide_passport(tmp_path):
    """Просроченный сертификат не вытесняет паспорт — в КП идёт паспорт."""

    async def scenario():
        async with open_session(tmp_path) as db:
            version = await make_version(db, supplier_name="Supplier A")
            item = await add_item(db, version, "Огнетушитель ОП-5(з)", article="ОП-5")
            await add_price(db, item, "1000.00")
            canonical = CanonicalProduct(
                match_key=f"test-{uuid.uuid4()}",
                title=item.source_name,
                normalized_title=item.normalized_name,
                normalized_article=item.normalized_article,
            )
            passport = ProductDocument(
                kind="passport",
                title="Паспорт ОП-5",
                storage_key="/files/passport-op5.pdf",
                source="passport_extraction",
            )
            expired = ProductDocument(
                kind="certificate",
                title="Сертификат ОП-5 (срок истёк)",
                registration_number="ЕАЭС RU С-RU.ПБ01.В.00003/26",
                valid_until=date(2024, 1, 1),
                storage_key="https://documents.example/cert-expired.pdf",
            )
            db.add_all([canonical, passport, expired])
            await db.flush()
            db.add(
                DocumentLink(
                    document_id=expired.id,
                    canonical_product_id=canonical.id,
                    confirmed_by="manager@example.test",
                    evidence={"verdict": "valid", "valid_until": "2024-01-01"},
                )
            )
            db.add(
                DocumentLink(
                    document_id=passport.id,
                    canonical_product_id=canonical.id,
                    confirmed_by="demo_pipeline_wave2",
                )
            )
            await db.flush()

            calculation = await CalculationService(db).create(
                project_id=PROJECT_ID,
                lines=[
                    LineInput(
                        price_item_id=item.id,
                        quantity=Decimal("1"),
                        canonical_id=canonical.id,
                    )
                ],
            )
            row = calculation.items[0]
            assert row.certificate_number == passport.title
            assert row.certificate_reference == passport.storage_key
            assert row.certificate_confirmed_by == "demo_pipeline_wave2"

    asyncio.run(scenario())


def test_volume_discount_wins_when_order_is_large(tmp_path):
    """Объёмная скидка сравнивается с уровневой по проценту, а не заменяет её."""

    async def scenario():
        async with open_session(tmp_path) as db:
            version = await make_version(db, supplier_name="Supplier B")
            item = await add_item(db, version, "Шкаф пожарный ШПК-320", article="ШПК-320")
            await add_price(db, item, "10000.00")
            db.add(rule(version.supplier_id, DiscountScope.supplier.value, "2"))
            db.add(
                rule(
                    version.supplier_id,
                    DiscountScope.volume.value,
                    "4",
                    min_amount=Decimal("300000"),
                )
            )
            await db.flush()

            service = CalculationService(db)
            # 40 шт. по 10 000 со скидкой 2% — 392 000 ₽, порог пройден.
            calculation = await service.create(
                project_id=PROJECT_ID,
                lines=[LineInput(price_item_id=item.id, quantity=Decimal("40"))],
            )
            row = calculation.items[0]
            assert row.discount_source == DiscountScope.volume.value
            assert row.discount_percent == Decimal("4")

    asyncio.run(scenario())


def test_small_order_keeps_level_discount(tmp_path):
    async def scenario():
        async with open_session(tmp_path) as db:
            version = await make_version(db, supplier_name="Supplier B")
            item = await add_item(db, version, "Шкаф пожарный ШПК-320", article="ШПК-320")
            await add_price(db, item, "10000.00")
            db.add(rule(version.supplier_id, DiscountScope.supplier.value, "2"))
            db.add(
                rule(
                    version.supplier_id,
                    DiscountScope.volume.value,
                    "4",
                    min_amount=Decimal("300000"),
                )
            )
            await db.flush()

            calculation = await CalculationService(db).create(
                project_id=PROJECT_ID,
                lines=[LineInput(price_item_id=item.id, quantity=Decimal("2"))],
            )
            row = calculation.items[0]
            assert row.discount_source == DiscountScope.supplier.value
            assert row.price_after_discount == Decimal("9800.00")

    asyncio.run(scenario())


def test_category_discount_uses_manual_category(tmp_path):
    async def scenario():
        async with open_session(tmp_path) as db:
            version = await make_version(db, supplier_name="Supplier A")
            category = GlobalCategory(
                project_id=PROJECT_ID,
                name="Огнетушители",
                normalized_name=f"ОГНЕТУШИТЕЛИ-{uuid.uuid4().hex[:6]}",
                path="Огнетушители",
                normalized_path=f"ОГНЕТУШИТЕЛИ-{uuid.uuid4().hex[:6]}",
            )
            db.add(category)
            await db.flush()
            item = await add_item(db, version, "Огнетушитель ОП-4(з)", article="ОП-4")
            item.manual_category_id = category.id
            await add_price(db, item, "500.00")
            db.add(
                rule(
                    version.supplier_id,
                    DiscountScope.category.value,
                    "24",
                    category_id=category.id,
                )
            )
            await db.flush()

            calculation = await CalculationService(db).create(
                project_id=PROJECT_ID,
                lines=[LineInput(price_item_id=item.id, quantity=Decimal("1"))],
            )
            row = calculation.items[0]
            assert row.discount_source == DiscountScope.category.value
            assert row.price_after_discount == Decimal("380.00")

    asyncio.run(scenario())


def test_quantity_tier_is_picked_before_discounts(tmp_path):
    """Ступень «от 50 шт.» — цена прайса, а не скидка; она применяется первой."""

    async def scenario():
        async with open_session(tmp_path) as db:
            version = await make_version(db, supplier_name="Supplier A")
            item = await add_item(db, version, "Рукав пожарный РПК-50", article="РПК-50")
            await add_price(db, item, "1000.00", min_quantity="1")
            await add_price(db, item, "800.00", min_quantity="50")
            db.add(rule(version.supplier_id, DiscountScope.supplier.value, "10"))
            await db.flush()

            calculation = await CalculationService(db).create(
                project_id=PROJECT_ID,
                lines=[LineInput(price_item_id=item.id, quantity=Decimal("60"))],
            )
            row = calculation.items[0]
            assert row.base_price == Decimal("800.00")
            assert row.price_after_discount == Decimal("720.00")

    asyncio.run(scenario())


# ------------------------------------------------------------------ ручная правка


def test_manual_price_survives_recalculation(tmp_path):
    """Пересчёт скидок не должен затирать согласованную с клиентом цену."""

    async def scenario():
        async with open_session(tmp_path) as db:
            version = await make_version(db, supplier_name="Supplier A")
            item = await add_item(db, version, "Огнетушитель ОП-4(з)", article="ОП-4")
            await add_price(db, item, "1000.00")
            db.add(rule(version.supplier_id, DiscountScope.supplier.value, "10"))
            await db.flush()

            service = CalculationService(db)
            calculation = await service.create(
                project_id=PROJECT_ID,
                lines=[LineInput(price_item_id=item.id, quantity=Decimal("2"))],
            )
            row = calculation.items[0]
            await service.override_price(
                row.id, price=Decimal("850.00"), reason="согласовано с клиентом", changed_by="ivanov"
            )
            totals = await service.recalculate(calculation)
            assert row.manual_price == Decimal("850.00")
            assert row.price_after_discount == Decimal("900.00")
            assert totals.final == Decimal("1700.00")
            assert totals.discounted == Decimal("1800.00")

    asyncio.run(scenario())


def test_override_without_reason_is_rejected(tmp_path):
    async def scenario():
        async with open_session(tmp_path) as db:
            version = await make_version(db, supplier_name="Supplier A")
            item = await add_item(db, version, "Огнетушитель ОП-4(з)")
            await add_price(db, item, "100.00")
            service = CalculationService(db)
            calculation = await service.create(
                project_id=PROJECT_ID, lines=[LineInput(price_item_id=item.id)]
            )
            with pytest.raises(ValueError):
                await service.override_price(
                    calculation.items[0].id, price=Decimal("90"), reason="   ", changed_by="ivanov"
                )

    asyncio.run(scenario())


def test_every_override_lands_in_the_log(tmp_path):
    """Журнал append-only: видно всю цепочку правок, а не последнюю цену."""

    async def scenario():
        async with open_session(tmp_path) as db:
            version = await make_version(db, supplier_name="Supplier A")
            item = await add_item(db, version, "Огнетушитель ОП-4(з)")
            await add_price(db, item, "1000.00")
            service = CalculationService(db)
            calculation = await service.create(
                project_id=PROJECT_ID, lines=[LineInput(price_item_id=item.id)]
            )
            row_id = calculation.items[0].id
            await service.override_price(
                row_id, price=Decimal("900"), reason="скидка по договорённости", changed_by="ivanov"
            )
            await service.override_price(
                row_id, price=Decimal("880"), reason="торг на встрече", changed_by="petrov"
            )
            history = await service.history(row_id)
            assert [entry.new_price for entry in history] == [Decimal("900"), Decimal("880")]
            assert history[1].previous_price == Decimal("900")
            assert [entry.changed_by for entry in history] == ["ivanov", "petrov"]

    asyncio.run(scenario())


def test_calculation_with_unknown_vat_mode_keeps_price_unchanged(tmp_path):
    """P0-2 на уровне расчёта: vat_mode=unknown со ставкой не завышает КП."""

    async def scenario():
        async with open_session(tmp_path) as db:
            version = await make_version(db, supplier_name="Supplier A")
            item = await add_item(db, version, "Кран пожарный КП")
            value = await add_price(db, item, "618.40")
            value.vat_rate_percent = Decimal("22")
            await db.flush()

            calculation = await CalculationService(db).create(
                project_id=PROJECT_ID,
                lines=[LineInput(price_item_id=item.id, quantity=Decimal("1"))],
            )
            row = calculation.items[0]
            assert row.vat_mode == "unknown"
            assert (row.price_without_vat, row.vat_amount, row.price_with_vat) == (
                Decimal("618.40"),
                None,
                Decimal("618.40"),
            )
            assert calculation.total_amount == Decimal("618.40")

    asyncio.run(scenario())
