"""КП из заявки: заявка → подбор (match-поля) → quote → request_id → Excel.

Поведенческий цикл из спецификации issue:
- заявка с текстом → разобранные строки (PriceRequestLine) → match-поля заполнены;
- POST /orders/{id}/quote собирает LineInput по канону строки и первому офферу
  активного прайса, создаёт Calculation с request_id == id заявки;
- строки без цены в активных прайсах пропускаются (skipped), а не падают 500;
- подтверждённый DocumentLink даёт certificate_number в позиции расчёта;
- export.xlsx (build_quote_workbook) собирается и содержит ссылку на документ.
"""

from __future__ import annotations

import asyncio
import io
import uuid
from datetime import date
from decimal import Decimal

import pytest
from fastapi import HTTPException
from openpyxl import load_workbook
from sqlalchemy import select

from app.api.calculations import get_calculation
from app.api.orders import (
    CreateOrderIn,
    QuoteSelection,
    QuoteSelectionsIn,
    create_order,
    create_quote_from_order,
    get_line_offers,
)
from app.matching.canonical import CanonicalService
from app.matching.export import build_quote_workbook
from app.matching.text import article_key
from app.models import (
    Calculation,
    CanonicalProduct,
    DocumentLink,
    PriceRequestLine,
    ProductDocument,
    Project,
    SupplierPriceValue,
)
from tests.test_matching import PROJECT_ID, add_item, make_version, open_session

CERT_NUMBER = "ЕАЭС RU С-RU.ПБ01.В.00001/26"
CERT_REFERENCE = "https://documents.example/cert-op5.pdf"
RAW_TEXT = "\n".join(
    [
        "Огнетушитель ОП-5(з) — 2 шт",
        "Огнетушитель ОП-4 — 3 шт",
        "Огнетушитель ОП-99 — 5 шт",
    ]
)


async def add_price(db, item, amount: str) -> SupplierPriceValue:
    value = SupplierPriceValue(
        item_id=item.id,
        price_kind="dealer",
        amount=Decimal(amount),
        currency="RUB",
    )
    db.add(value)
    await db.flush()
    return value


def test_quote_from_request_fills_request_id_items_and_export(tmp_path) -> None:
    """Заявка с подобранными канонами → КП: request_id, позиции, сертификат, xlsx."""

    async def scenario() -> None:
        async with open_session(tmp_path) as db:
            version = await make_version(db, supplier_name="Supplier A")
            item5 = await add_item(
                db, version, "Огнетушитель ОП-5(з)-АВСЕ", article="ОП-5"
            )
            await add_price(db, item5, "1000.00")
            item4 = await add_item(db, version, "Огнетушитель ОП-4(з)", article="ОП-4")
            await add_price(db, item4, "850.00")

            rebuild = await CanonicalService(db).rebuild()
            assert rebuild.items == 2
            assert rebuild.canonicals == 2

            canon_by_article = {
                row.normalized_article: row
                for row in (await db.scalars(select(CanonicalProduct)))
            }
            canon5 = canon_by_article[article_key("ОП-5")]
            canon4 = canon_by_article[article_key("ОП-4")]

            # Канон без источников/цены: строка с ним должна уйти в skipped.
            canon99 = CanonicalProduct(
                match_key=f"no-price-{uuid.uuid4().hex}",
                title="Огнетушитель ОП-99(з)",
                normalized_title="ОГНЕТУШИТЕЛЬ ОП-99(З)",
                normalized_article=article_key("ОП-99"),
            )
            db.add(canon99)
            await db.flush()

            # Подтверждённый сертификат на канон ОП-5.
            certificate = ProductDocument(
                kind="certificate",
                title="Сертификат ОП-5",
                registration_number=CERT_NUMBER,
                valid_until=date(2028, 1, 1),
                storage_key=CERT_REFERENCE,
            )
            db.add(certificate)
            await db.flush()
            db.add(
                DocumentLink(
                    document_id=certificate.id,
                    canonical_product_id=canon5.id,
                    confirmed_by="manager@example.test",
                    evidence={"verdict": "valid", "valid_until": "2028-01-01"},
                )
            )
            await db.flush()

            created = await create_order(
                payload=CreateOrderIn(
                    title="Заявка ОП-4/ОП-5",
                    client="\u0414\u0435\u043C\u043E\u0421\u0442\u0440\u043E\u0439",
                    raw_text=RAW_TEXT,
                ),
                session=db,
                user=None,
            )
            order_id = uuid.UUID(created["id"])
            by_query = {row["query_name"]: row for row in created["lines"]}
            line5 = uuid.UUID(by_query["Огнетушитель ОП-5(з)"]["id"])
            line4 = uuid.UUID(by_query["Огнетушитель ОП-4"]["id"])
            line99 = uuid.UUID(by_query["Огнетушитель ОП-99"]["id"])

            # Подбор (этап issues 1–2) закрепил каноны в строках.
            line5_row = await db.get(PriceRequestLine, line5)
            line4_row = await db.get(PriceRequestLine, line4)
            line99_row = await db.get(PriceRequestLine, line99)
            line5_row.matched_canonical_id = canon5.id
            line5_row.match_stage = "auto"
            line5_row.match_score = 0.95
            line4_row.matched_canonical_id = canon4.id
            line4_row.match_stage = "auto"
            line4_row.match_score = 0.9
            line99_row.matched_canonical_id = canon99.id
            line99_row.match_stage = "manual"
            await db.commit()

            project = await db.get(Project, PROJECT_ID)
            assert project is not None

            response = await create_quote_from_order(
                order_id=order_id,
                session=db,
                project=project,
                user=None,
            )

            assert response["request_id"] == str(order_id)
            assert response["title"] == "Заявка ОП-4/ОП-5"
            assert len(response["items"]) == 2
            assert [item["ordinal"] for item in response["items"]] == [1, 2]
            assert len(response["skipped"]) == 1
            skipped = response["skipped"][0]
            assert skipped["line_id"] == str(line99)
            assert "нет цены в активных прайсах" in skipped["reason"]

            calc_id = uuid.UUID(response["id"])
            calculation = await db.get(Calculation, calc_id)
            assert calculation is not None
            assert calculation.request_id == order_id
            assert len(calculation.items) == 2

            item5_calc = next(
                item for item in calculation.items if item.source_line == RAW_TEXT.splitlines()[0]
            )
            assert item5_calc.certificate_number == CERT_NUMBER
            assert item5_calc.certificate_confirmed_by == "manager@example.test"
            assert item5_calc.certificate_reference == CERT_REFERENCE
            assert item5_calc.quantity == Decimal("2")

            # GET /calculations/{id} отдаёт request_id.
            card = await get_calculation(
                calculation_id=calc_id,
                session=db,
                project=project,
                _reader=None,
            )
            assert card["request_id"] == str(order_id)

            # Экспорт КП собирается: блок тендера (ID/заказчик, 47a86fe7) идёт
            # до строки заголовков — ищем её по «Наименование», а не по sheet[1].
            workbook = load_workbook(io.BytesIO(build_quote_workbook(calculation)))
            sheet = workbook.active
            assert sheet["B1"].value == "ID тендера"
            assert sheet["B2"].value == "Заказчик тендера"
            assert sheet["C2"].value == "\u0414\u0435\u043C\u043E\u0421\u0442\u0440\u043E\u0439"
            header_row = next(
                (
                    row
                    for row in sheet.iter_rows(max_row=6)
                    if "Наименование" in [cell.value for cell in row]
                ),
                None,
            )
            assert header_row is not None
            headings = [cell.value for cell in header_row]
            assert "Ссылка/файл сертификата" in headings
            reference_col = headings.index("Ссылка/файл сертификата") + 1
            found = False
            for row in sheet.iter_rows(min_row=header_row[0].row + 1, values_only=False):
                if row[1].value == RAW_TEXT.splitlines()[0]:
                    reference_cell = row[reference_col - 1]
                    assert reference_cell.value == CERT_REFERENCE
                    assert reference_cell.hyperlink.target == CERT_REFERENCE
                    found = True
            assert found
            workbook.close()

    asyncio.run(scenario())


def test_quote_without_prices_skips_lines_and_returns_card(tmp_path) -> None:
    """Строки без цены в активных прайсах не роняют quote: пустой расчёт + skipped."""

    async def scenario() -> None:
        async with open_session(tmp_path) as db:
            version = await make_version(db, supplier_name="Supplier A")
            item = await add_item(db, version, "Огнетушитель ОП-5(з)", article="ОП-5")
            await add_price(db, item, "1000.00")
            await CanonicalService(db).rebuild()
            canon5 = (
                await db.scalars(
                    select(CanonicalProduct).where(CanonicalProduct.normalized_article == article_key("ОП-5"))
                )
            ).one()

            created = await create_order(
                payload=CreateOrderIn(
                    client="\u0414\u0435\u043C\u043E\u0421\u0442\u0440\u043E\u0439",
                    raw_text="Огнетушитель ОП-99 — 5 шт",
                ),
                session=db,
                user=None,
            )
            order_id = uuid.UUID(created["id"])
            line_row = await db.get(
                PriceRequestLine, uuid.UUID(created["lines"][0]["id"])
            )
            # Подбор не нашёл цены — строка осталась без канона.
            line_row.matched_canonical_id = None
            line_row.match_stage = "none"
            await db.commit()

            project = await db.get(Project, PROJECT_ID)
            response = await create_quote_from_order(
                order_id=order_id,
                session=db,
                project=project,
                user=None,
            )
            assert response["items"] == []
            assert response["request_id"] == str(order_id)
            assert len(response["skipped"]) == 1
            assert response["skipped"][0]["reason"] == "строка не подобрана: нет канона"

            # Явный выбор позиции менеджером всё равно позволяет собрать КП.
            override = await create_quote_from_order(
                order_id=order_id,
                payload=QuoteSelectionsIn(
                    selections=[
                        QuoteSelection(
                            line_id=uuid.UUID(created["lines"][0]["id"]),
                            canonical_id=canon5.id,
                            price_item_id=item.id,
                            quantity=4,
                        )
                    ]
                ),
                session=db,
                project=project,
                user=None,
            )
            assert len(override["items"]) == 1
            assert override["items"][0]["quantity"] == 4.0
            assert override["skipped"] == []

    asyncio.run(scenario())


def test_line_offers_found_by_name_without_designation(tmp_path) -> None:
    """Задача №7, дополнение: пикер ищет предложения по имени позиции.

    Автоподбор вкладки требует дефисного обозначения («ОП-5»); строка без
    него должна всё равно находить канон — каскадом (fuzzy), а не тишиной.
    """

    async def scenario() -> None:
        async with open_session(tmp_path) as db:
            version = await make_version(db, supplier_name="Supplier A")
            item = await add_item(db, version, "Гидроклапан предохранительный МКПВ20/3С2")
            await add_price(db, item, "1044.80")
            rebuild = await CanonicalService(db).rebuild()
            assert rebuild.canonicals == 1

            created = await create_order(
                payload=CreateOrderIn(
                    client="\u0414\u0435\u043C\u043E\u0421\u0442\u0440\u043E\u0439",
                    raw_text="Гидроклапан предохранительный МКПВ20/3С2 — 1 шт",
                ),
                session=db,
                user=None,
            )
            order_id = uuid.UUID(created["id"])
            line_id = uuid.UUID(created["lines"][0]["id"])

            response = await get_line_offers(
                order_id=order_id, line_id=line_id, user=None, session=db
            )
            assert response["lineId"] == str(line_id)
            assert response["query"] == "Гидроклапан предохранительный МКПВ20/3С2"
            assert response["offers"], "предложения должны найтись по имени"
            offer = response["offers"][0]
            assert offer["priceItemId"] == str(item.id)
            assert offer["supplier"] == "Supplier A"
            assert offer["price"] == 1044.8
            assert offer["canonicalId"]
            assert offer["canonicalTitle"]

            # Чужая строка — 404, а не пустой список чужих предложений.
            with pytest.raises(HTTPException) as foreign:
                await get_line_offers(
                    order_id=order_id,
                    line_id=uuid.uuid4(),
                    user=None,
                    session=db,
                )
            assert foreign.value.status_code == 404

    asyncio.run(scenario())


def test_line_offers_empty_when_nothing_matches(tmp_path) -> None:
    """Несуществующий товар: пикер честно пуст, а не падает и не выдумывает."""

    async def scenario() -> None:
        async with open_session(tmp_path) as db:
            await make_version(db, supplier_name="Supplier A")

            created = await create_order(
                payload=CreateOrderIn(client="\u0414\u0435\u043C\u043E\u0421\u0442\u0440\u043E\u0439", raw_text="Лунный модуль — 1 шт"),
                session=db,
                user=None,
            )

            response = await get_line_offers(
                order_id=uuid.UUID(created["id"]),
                line_id=uuid.UUID(created["lines"][0]["id"]),
                user=None,
                session=db,
            )
            assert response["offers"] == []

    asyncio.run(scenario())