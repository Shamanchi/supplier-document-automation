"""P0 contracts for the asynchronous external price-import service."""

from __future__ import annotations

import asyncio
import io
import uuid
from decimal import Decimal

import openpyxl
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

from app.currencies import normalise_currency_hint
from app.db import get_session
from app.import_auth import api_key_digest, generate_service_api_key
from app.import_service import PROGRESS, transition
from app.import_service import advance_price_import
from app.import_telemetry import aggregate_import_telemetry
from app.models import (
    Base,
    Currency,
    PriceImport,
    PriceImportStatus,
    PriceFile,
    PriceParseIssue,
    PriceVersion,
    RuleGenerationAttempt,
    ServiceClient,
    Supplier,
    SupplierPriceItem,
    SupplierPriceValue,
    VersionStatus,
)
from app.parsing.grid import Cell, SheetGrid
from app.parsing.mapping import ColumnMapping, MappingDecision, TableMapping
from app.parsing.pipeline import parse_grids
from app.webhook_service import (
    CallbackDestination,
    _post_pinned_https,
    sign_webhook_payload,
    validate_callback_url,
)
from app.reconciliation import reconcile_version


def test_completeness_failure_downgrades_validated_version(monkeypatch) -> None:
    asyncio.run(_quality_status_case(monkeypatch, completeness_passed=False, errors_count=0))


def test_blocking_parse_errors_downgrade_validated_version(monkeypatch) -> None:
    asyncio.run(_quality_status_case(monkeypatch, completeness_passed=True, errors_count=1))


def test_complete_deterministic_result_remains_validated(monkeypatch) -> None:
    asyncio.run(
        _quality_status_case(
            monkeypatch,
            completeness_passed=True,
            errors_count=0,
            expected_import_status=PriceImportStatus.completed.value,
            expected_version_status=VersionStatus.validated.value,
        )
    )


async def _quality_status_case(
    monkeypatch,
    *,
    completeness_passed: bool,
    errors_count: int,
    expected_import_status: str = PriceImportStatus.review_required.value,
    expected_version_status: str = VersionStatus.validation_failed.value,
) -> None:
    import app.import_service as import_module

    async def money_gate(*_args, **_kwargs):
        return {"passed": True, "failures": 0}

    async def completeness_gate(*_args, **_kwargs):
        return {
            "passed": completeness_passed,
            "confirmed_missed_positions": 0 if completeness_passed else 1,
        }

    async def category_gate(*_args, **_kwargs):
        return {"passed": True, "failures": 0}

    monkeypatch.setattr(import_module, "apply_money_gate", money_gate)
    monkeypatch.setattr(import_module, "reconcile_version", completeness_gate)
    monkeypatch.setattr(import_module, "apply_category_context_gate", category_gate)
    monkeypatch.setattr(import_module, "aggregate_import_telemetry", lambda *_args, **_kwargs: _empty_telemetry())

    engine = create_async_engine("sqlite+aiosqlite://")
    factory = async_sessionmaker(engine, expire_on_commit=False)
    async with engine.begin() as connection:
        await connection.run_sync(Base.metadata.create_all)
    try:
        async with factory() as session:
            client = ServiceClient(name="P0 quality", key_prefix="p0q", key_digest="q" * 64)
            supplier = Supplier(code=f"p0-{uuid.uuid4()}", name="P0")
            session.add_all([client, supplier])
            await session.flush()
            price_file = PriceFile(
                supplier_id=supplier.id,
                original_filename="quality.xlsx",
                storage_key="quality-file",
                sha256="q" * 64,
                size_bytes=1,
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
            session.add(price_file)
            await session.flush()
            version = PriceVersion(
                supplier_id=supplier.id,
                price_file_id=price_file.id,
                version_no=1,
                status=VersionStatus.validated.value,
                item_count=1,
                rows_accepted=1,
                errors_count=errors_count,
            )
            session.add(version)
            await session.flush()
            price_import = PriceImport(
                service_client_id=client.id,
                supplier_id=supplier.id,
                price_file_id=price_file.id,
                price_version_id=version.id,
                idempotency_key=f"quality-{uuid.uuid4()}",
                request_fingerprint="q" * 64,
                vat_rate_percent=Decimal("22"),
                status=PriceImportStatus.parsing.value,
                stage=PriceImportStatus.parsing.value,
            )
            session.add(price_import)
            await session.commit()

            status = await import_module._quality_gate(
                session,
                price_import=price_import,
                version=version,
                price_file=price_file,
                data=b"x",
            )

            assert status == expected_import_status
            assert version.status == expected_version_status
    finally:
        await engine.dispose()


async def _empty_telemetry() -> dict:
    return {
        "operations": 0,
        "input_tokens": 0,
        "output_tokens": 0,
        "total_tokens": 0,
        "cached_input_tokens": 0,
        "billable_input_tokens": 0,
        "http_retries": 0,
        "semantic_retries": 0,
        "llm_latency_ms": 0,
        "cache_hit_ratio": 0.0,
    }


def test_currency_hint_accepts_iso_and_rejects_arbitrary_three_letters() -> None:
    assert normalise_currency_hint(" kzt ") == "KZT"
    assert normalise_currency_hint(None) is None
    with pytest.raises(ValueError):
        normalise_currency_hint("ZZZ")


def test_service_key_is_hashed_and_raw_secret_is_not_persisted(monkeypatch) -> None:
    from app.config import settings

    monkeypatch.setattr(settings, "service_api_key_pepper", "p" * 48)
    raw, prefix, digest = generate_service_api_key()
    assert raw.startswith(f"plk_{prefix}_")
    assert digest == api_key_digest(raw)
    assert raw not in digest


def test_terminal_transition_has_documented_progress() -> None:
    price_import = PriceImport(status="queued", stage="queued", progress_percent=0)
    price_import.id = __import__("uuid").uuid4()
    event = transition(price_import, PriceImportStatus.review_required.value)
    assert price_import.status == "review_required"
    assert price_import.progress_percent == PROGRESS["review_required"] == 100
    assert event.from_status == "queued"
    assert event.to_status == "review_required"
    assert price_import.finished_at is not None


def test_webhook_signature_is_stable_and_private_destinations_are_rejected(monkeypatch) -> None:
    assert sign_webhook_payload("secret", "123", b"{}") == sign_webhook_payload(
        "secret", "123", b"{}"
    )
    monkeypatch.setattr(
        "app.webhook_service.socket.getaddrinfo",
        lambda *_args, **_kwargs: [(2, 1, 6, "", ("127.0.0.1", 443))],
    )
    with pytest.raises(ValueError):
        validate_callback_url("https://localhost/hook")


def test_webhook_connects_to_verified_ip_with_original_tls_identity(monkeypatch) -> None:
    calls: list[tuple] = []

    class Reader:
        def __init__(self) -> None:
            self.lines = iter([b"HTTP/1.1 204 No Content\r\n", b"\r\n"])

        async def readline(self) -> bytes:
            return next(self.lines, b"")

    class Writer:
        def write(self, data: bytes) -> None:
            calls.append(("request", data))

        async def drain(self) -> None:
            return None

        def close(self) -> None:
            return None

        async def wait_closed(self) -> None:
            return None

    async def open_connection(host, port, **kwargs):
        calls.append((host, port, kwargs.get("server_hostname")))
        return Reader(), Writer()

    monkeypatch.setattr(
        "app.webhook_service.resolve_callback_destination",
        lambda _url: CallbackDestination(
            "https://hooks.example.test/path",
            "hooks.example.test",
            443,
            ("93.184.216.34",),
            "/path",
        ),
    )
    monkeypatch.setattr("app.webhook_service.asyncio.open_connection", open_connection)
    monkeypatch.setattr("app.webhook_service.ssl.create_default_context", lambda: object())
    status = asyncio.run(
        _post_pinned_https(
            "https://hooks.example.test/path",
            b"{}",
            {"Content-Type": "application/json"},
        )
    )
    assert status == 204
    assert calls[0] == ("93.184.216.34", 443, "hooks.example.test")
    assert b"Host: hooks.example.test\r\n" in calls[1][1]


def test_horizontal_panels_keep_distinct_provenance_and_money_semantics() -> None:
    grid = SheetGrid(
        name="Прайс",
        rows=[
            [
                Cell(value="Наименование"),
                Cell(value="Цена"),
                Cell(value="Наименование"),
                Cell(value="Цена"),
            ],
            [
                Cell(value="Товар A"),
                Cell(value="130-300"),
                Cell(value="Товар B"),
                Cell(value="50 тенге за 1 км"),
            ],
        ],
    )
    decision = MappingDecision(
        auto_ok=True,
        tables=[
            TableMapping(
                sheet_name="Прайс",
                header_row_index=0,
                source_section_id="Прайс:1",
                allow_qualified_prices=True,
                columns=[
                    ColumnMapping(index=0, role="name"),
                    ColumnMapping(index=1, role="price"),
                    ColumnMapping(index=2, role="name"),
                    ColumnMapping(index=3, role="price"),
                ],
            )
        ],
    )
    parsed = parse_grids([grid], decision)
    assert [item.source_name for item in parsed.items] == ["Товар A", "Товар B"]
    assert [item.source_panel_index for item in parsed.items] == [0, 1]
    assert parsed.items[0].values[0].qualifier == "range"
    assert parsed.items[0].values[0].amount_from == Decimal("130.00")
    assert parsed.items[0].values[0].amount_to == Decimal("300.00")
    assert parsed.items[1].values[0].currency == "KZT"
    assert parsed.items[1].values[0].basis_unit == "км"


def test_multiple_price_tiers_require_and_preserve_explicit_threshold_links() -> None:
    grid = SheetGrid(
        name="Прайс",
        rows=[
            [
                Cell(value="Наименование"),
                Cell(value="От, шт"),
                Cell(value="Цена 1"),
                Cell(value="От, шт"),
                Cell(value="Цена 2"),
            ],
            [
                Cell(value="Товар"),
                Cell(value=10),
                Cell(value=100),
                Cell(value=100),
                Cell(value=80),
            ],
        ],
    )

    def decision(*, linked: bool) -> MappingDecision:
        return MappingDecision(
            auto_ok=True,
            tables=[
                TableMapping(
                    sheet_name="Прайс",
                    header_row_index=0,
                    columns=[
                        ColumnMapping(index=0, role="name"),
                        ColumnMapping(index=1, role="quantity_threshold"),
                        ColumnMapping(
                            index=2,
                            role="price",
                            quantity_threshold_column_index=1 if linked else None,
                        ),
                        ColumnMapping(index=3, role="quantity_threshold"),
                        ColumnMapping(
                            index=4,
                            role="price",
                            quantity_threshold_column_index=3 if linked else None,
                        ),
                    ],
                )
            ],
        )

    parsed = parse_grids([grid], decision(linked=True))
    assert [value.min_quantity for value in parsed.items[0].values] == [
        Decimal("10.000"),
        Decimal("100.000"),
    ]
    ambiguous = parse_grids([grid], decision(linked=False))
    assert any(issue.code == "PRICE_TIER_AMBIGUOUS" for issue in ambiguous.issues)


def test_conflicting_cell_and_header_currency_is_blocking_evidence() -> None:
    grid = SheetGrid(
        name="Прайс",
        rows=[
            [Cell(value="Наименование"), Cell(value="Цена, тенге")],
            [Cell(value="Товар"), Cell(value="100 USD")],
        ],
    )
    parsed = parse_grids(
        [grid],
        MappingDecision(
            auto_ok=True,
            tables=[
                TableMapping(
                    sheet_name="Прайс",
                    header_row_index=0,
                    columns=[
                        ColumnMapping(index=0, role="name"),
                        ColumnMapping(index=1, role="price", header="Цена, тенге"),
                    ],
                )
            ],
        ),
    )
    assert parsed.items[0].values[0].currency == "USD"
    assert any(issue.code == "PRICE_CURRENCY_CONFLICT" for issue in parsed.issues)


def test_database_enforces_client_scoped_idempotency_and_supplier_identity() -> None:
    asyncio.run(_database_constraints())


async def _database_constraints() -> None:
    engine = create_async_engine("sqlite+aiosqlite://")
    factory = async_sessionmaker(engine, expire_on_commit=False)
    async with engine.begin() as connection:
        await connection.run_sync(Base.metadata.create_all)
    try:
        async with factory() as session:
            clients = [
                ServiceClient(name="A", key_prefix="a", key_digest="a" * 64),
                ServiceClient(name="B", key_prefix="b", key_digest="b" * 64),
            ]
            session.add_all(clients)
            await session.flush()
            suppliers = [
                Supplier(
                    service_client_id=client.id,
                    external_id="supplier-1",
                    code=f"supplier-{index}",
                    name="Supplier",
                )
                for index, client in enumerate(clients)
            ]
            session.add_all(suppliers)
            await session.flush()
            for index, supplier in enumerate(suppliers):
                price_file = PriceFile(
                    supplier_id=supplier.id,
                    original_filename="price.xlsx",
                    storage_key=f"file-{index}",
                    sha256=str(index) * 64,
                    size_bytes=1,
                    mime="application/octet-stream",
                )
                session.add(price_file)
                await session.flush()
                version = PriceVersion(
                    supplier_id=supplier.id,
                    price_file_id=price_file.id,
                    version_no=1,
                )
                session.add(version)
                await session.flush()
                session.add(
                    PriceImport(
                        service_client_id=clients[index].id,
                        supplier_id=supplier.id,
                        price_file_id=price_file.id,
                        price_version_id=version.id,
                        idempotency_key="same-key",
                        request_fingerprint=f"fingerprint-{index}",
                    )
                )
            await session.commit()
            assert len((await session.execute(select(PriceImport))).scalars().all()) == 2
    finally:
        await engine.dispose()


def test_import_telemetry_counts_failed_operation_usage() -> None:
    asyncio.run(_failed_operation_usage())


async def _failed_operation_usage() -> None:
    engine = create_async_engine("sqlite+aiosqlite://")
    factory = async_sessionmaker(engine, expire_on_commit=False)
    async with engine.begin() as connection:
        await connection.run_sync(Base.metadata.create_all)
    try:
        async with factory() as session:
            client = ServiceClient(name="T", key_prefix="t", key_digest="t" * 64)
            supplier = Supplier(code="telemetry", name="Telemetry")
            session.add_all([client, supplier])
            await session.flush()
            price_file = PriceFile(
                supplier_id=supplier.id,
                original_filename="price.xlsx",
                storage_key="telemetry-file",
                sha256="f" * 64,
                size_bytes=1,
                mime="application/octet-stream",
            )
            session.add(price_file)
            await session.flush()
            version = PriceVersion(
                supplier_id=supplier.id, price_file_id=price_file.id, version_no=1
            )
            session.add(version)
            await session.flush()
            price_import = PriceImport(
                service_client_id=client.id,
                supplier_id=supplier.id,
                price_file_id=price_file.id,
                price_version_id=version.id,
                idempotency_key="telemetry-key",
                request_fingerprint="f" * 64,
            )
            session.add(price_import)
            await session.flush()
            event = {
                "operation_id": "failed-op",
                "retry_count": 2,
                "duration_ms": 900,
                "usage": {
                    "input_tokens": 100,
                    "output_tokens": 20,
                    "total_tokens": 120,
                    "cached_input_tokens": 60,
                    "billable_input_tokens": 40,
                },
            }
            session.add(
                RuleGenerationAttempt(
                    price_version_id=version.id,
                    price_import_id=price_import.id,
                    batch_id=__import__("uuid").uuid4(),
                    section_id="section",
                    status="failed",
                    provider="deepseek",
                    model="deepseek-v4-flash",
                    request_settings={},
                    prompt_version="v1",
                    prompt_id="prompt",
                    source_file_sha256="f" * 64,
                    source_context_hash="c" * 64,
                    input_context={},
                    prompt="prompt",
                    validation={"llm_telemetry_events": [event]},
                )
            )
            await session.commit()
            aggregate = await aggregate_import_telemetry(session, price_import.id)
            assert aggregate["operations"] == 1
            assert aggregate["total_tokens"] == 120
            assert aggregate["billable_input_tokens"] == 40
            assert aggregate["http_retries"] == 2
            assert aggregate["llm_latency_ms"] == 900
    finally:
        await engine.dispose()


def test_reconciliation_detects_missing_right_panel_on_an_already_parsed_row() -> None:
    asyncio.run(_missing_right_panel())


async def _missing_right_panel() -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Прайс"
    sheet.append(["Наименование", "Цена", "Наименование", "Цена"])
    sheet.append(["Товар слева", 100, "Товар справа", 200])
    output = io.BytesIO()
    workbook.save(output)
    data = output.getvalue()
    engine = create_async_engine("sqlite+aiosqlite://")
    factory = async_sessionmaker(engine, expire_on_commit=False)
    async with engine.begin() as connection:
        await connection.run_sync(Base.metadata.create_all)
    try:
        async with factory() as session:
            supplier = Supplier(code="panels", name="Panels")
            session.add(supplier)
            await session.flush()
            price_file = PriceFile(
                supplier_id=supplier.id,
                original_filename="panels.xlsx",
                storage_key="panels-file",
                sha256="p" * 64,
                size_bytes=len(data),
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
            session.add(price_file)
            await session.flush()
            version = PriceVersion(
                supplier_id=supplier.id, price_file_id=price_file.id, version_no=1
            )
            session.add(version)
            await session.flush()
            session.add(
                SupplierPriceItem(
                    price_version_id=version.id,
                    source_name="Товар слева",
                    normalized_name="товар слева",
                    source_sheet="Прайс",
                    source_row=2,
                    source_panel_index=0,
                    source_column_start=0,
                    source_column_end=2,
                )
            )
            await session.commit()
            summary = await reconcile_version(
                session, version=version, data=data, filename="panels.xlsx"
            )
            assert summary["passed"] is False
            assert summary["confirmed_missed_physical_rows"] == 1
            assert summary["confirmed_missed_positions"] == 1
            assert summary["missed_coordinates"][0]["panel_index"] == 1
    finally:
        await engine.dispose()


def test_http_idempotency_replays_without_creating_a_second_job(tmp_path, monkeypatch) -> None:
    from app.config import settings
    from app.main import app
    import app.api.price_imports as api_module
    import app.tasks as tasks_module

    database_path = tmp_path / "service.sqlite"
    engine = create_async_engine(f"sqlite+aiosqlite:///{database_path}")
    factory = async_sessionmaker(engine, expire_on_commit=False)
    monkeypatch.setattr(settings, "price_import_service_enabled", True)
    monkeypatch.setattr(settings, "service_api_key_pepper", "i" * 48)
    monkeypatch.setattr(api_module.storage, "put_price_file", lambda *_args: "stored/key")
    monkeypatch.setattr(api_module.storage, "ensure_buckets", lambda: None)

    async def no_quotas(*_args, **_kwargs) -> None:
        return None

    class FakeTask:
        @staticmethod
        def delay(*_args, **_kwargs):
            return None

    monkeypatch.setattr(api_module, "enforce_service_quotas", no_quotas)
    monkeypatch.setattr(tasks_module, "run_price_import", FakeTask())

    async def session_override():
        async with factory() as session:
            yield session

    async def prepare() -> str:
        async with engine.begin() as connection:
            await connection.run_sync(Base.metadata.create_all)
        raw, prefix, digest = generate_service_api_key()
        async with factory() as session:
            session.add(ServiceClient(name="HTTP client", key_prefix=prefix, key_digest=digest))
            await session.commit()
        return raw

    raw_key = asyncio.run(prepare())
    app.dependency_overrides[get_session] = session_override
    try:
        client = TestClient(app)
        headers = {
            "Authorization": f"Bearer {raw_key}",
            "Idempotency-Key": "stable-key-001",
        }
        form = {
            "supplier_external_id": "supplier-http",
            "supplier_name": "HTTP Supplier",
            "currency_hint": "KZT",
        }
        upload = {"file": ("price.xlsx", b"source bytes", "application/octet-stream")}
        first = client.post("/api/v1/price-imports", headers=headers, data=form, files=upload)
        second = client.post("/api/v1/price-imports", headers=headers, data=form, files=upload)
        conflict = client.post(
            "/api/v1/price-imports",
            headers=headers,
            data={**form, "supplier_name": "Changed"},
            files=upload,
        )
        assert first.status_code == 202
        assert first.json()["replayed"] is False
        assert second.status_code == 202
        assert second.json()["replayed"] is True
        assert second.json()["import_id"] == first.json()["import_id"]
        assert conflict.status_code == 409
        assert (
            client.get(
                f"/api/v1/price-imports/{first.json()['import_id']}/items",
                headers={"Authorization": f"Bearer {raw_key}"},
            ).status_code
            == 409
        )

        async def count_imports() -> int:
            async with factory() as session:
                return len((await session.execute(select(PriceImport))).scalars().all())

        assert asyncio.run(count_imports()) == 1
    finally:
        app.dependency_overrides.clear()
        asyncio.run(engine.dispose())


def test_upload_with_default_vat_rate_is_accepted_and_persisted(tmp_path, monkeypatch) -> None:
    """The optional НДС form field must survive request fingerprinting."""

    from app.config import settings
    from app.main import app
    import app.api.price_imports as api_module
    import app.tasks as tasks_module

    database_path = tmp_path / "vat.sqlite"
    engine = create_async_engine(f"sqlite+aiosqlite:///{database_path}")
    factory = async_sessionmaker(engine, expire_on_commit=False)
    monkeypatch.setattr(settings, "price_import_service_enabled", True)
    monkeypatch.setattr(settings, "service_api_key_pepper", "v" * 48)
    monkeypatch.setattr(api_module.storage, "put_price_file", lambda *_args: "stored/key")
    monkeypatch.setattr(api_module.storage, "ensure_buckets", lambda: None)

    async def no_quotas(*_args, **_kwargs) -> None:
        return None

    class FakeTask:
        @staticmethod
        def delay(*_args, **_kwargs):
            return None

    monkeypatch.setattr(api_module, "enforce_service_quotas", no_quotas)
    monkeypatch.setattr(tasks_module, "run_price_import", FakeTask())

    async def session_override():
        async with factory() as session:
            yield session

    async def prepare() -> str:
        async with engine.begin() as connection:
            await connection.run_sync(Base.metadata.create_all)
        raw, prefix, digest = generate_service_api_key()
        async with factory() as session:
            session.add(ServiceClient(name="VAT client", key_prefix=prefix, key_digest=digest))
            await session.commit()
        return raw

    raw_key = asyncio.run(prepare())
    app.dependency_overrides[get_session] = session_override
    try:
        client = TestClient(app)
        upload = {"file": ("price.xlsx", b"source bytes", "application/octet-stream")}
        response = client.post(
            "/api/v1/price-imports",
            headers={"Authorization": f"Bearer {raw_key}", "Idempotency-Key": "vat-key-001"},
            data={
                "supplier_external_id": "supplier-vat",
                "supplier_name": "VAT Supplier",
                "currency_hint": "KZT",
                "vat_rate_percent": "22",
            },
            files=upload,
        )
        assert response.status_code == 202, response.text

        async def stored_rate() -> Decimal | None:
            async with factory() as session:
                return (await session.execute(select(PriceImport))).scalar_one().vat_rate_percent

        assert asyncio.run(stored_rate()) == Decimal("22")
    finally:
        app.dependency_overrides.clear()
        asyncio.run(engine.dispose())


def test_browser_ui_uses_only_import_api_without_exposing_service_key(tmp_path, monkeypatch) -> None:
    from app.config import settings
    from app.main import app
    import app.api.price_imports as api_module
    import app.tasks as tasks_module

    database_path = tmp_path / "browser-service.sqlite"
    engine = create_async_engine(f"sqlite+aiosqlite:///{database_path}")
    factory = async_sessionmaker(engine, expire_on_commit=False)
    for key, value in {
        "app_admin_username": "ui-admin",
        "app_admin_password": "ui-admin-password-long",
        "app_user_username": "ui-viewer",
        "app_user_password": "ui-viewer-password-long",
        "app_auth_secret": "ui-auth-secret-that-is-long-enough-123456789",
        "app_auth_secure_cookies": False,
        # External keys remain disabled; the authenticated browser scope is
        # deliberately independent of this rollout flag.
        "price_import_service_enabled": False,
    }.items():
        monkeypatch.setattr(settings, key, value)
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Прайс"
    sheet.append(["Артикул", "Название", "Цена"])
    sheet.append(["ФК-1", "Фильтр кассетный", 1250])
    sheet.append(["ФК-2", "Фильтр карманный", 2400])
    workbook.create_sheet("Условия").append(["Доставка", "5 дней"])
    workbook_buffer = io.BytesIO()
    workbook.save(workbook_buffer)
    source_bytes = workbook_buffer.getvalue()
    monkeypatch.setattr(api_module.storage, "put_price_file", lambda *_args: "stored/browser")
    monkeypatch.setattr(
        api_module.storage, "get_price_file", lambda _storage_key: source_bytes
    )

    async def no_quotas(*_args, **_kwargs) -> None:
        return None

    class FakeTask:
        @staticmethod
        def delay(*_args, **_kwargs):
            return None

    monkeypatch.setattr(api_module, "enforce_service_quotas", no_quotas)
    monkeypatch.setattr(tasks_module, "run_price_import", FakeTask())

    async def session_override():
        async with factory() as session:
            yield session

    async def prepare() -> None:
        async with engine.begin() as connection:
            await connection.run_sync(Base.metadata.create_all)

    asyncio.run(prepare())
    app.dependency_overrides[get_session] = session_override
    try:
        client = TestClient(app)
        admin_session = client.post(
            "/api/v1/auth/login",
            json={"username": "ui-admin", "password": "ui-admin-password-long"},
        ).json()
        viewer_session = client.post(
            "/api/v1/auth/login",
            json={"username": "ui-viewer", "password": "ui-viewer-password-long"},
        ).json()
        admin_headers = {
            "Authorization": f"Bearer {admin_session['access_token']}",
            "Idempotency-Key": "browser-key-001",
        }
        viewer_headers = {"Authorization": f"Bearer {viewer_session['access_token']}"}
        form = {
            "supplier_external_id": "browser-supplier",
            "supplier_name": "Browser Supplier",
            "currency_hint": "KZT",
        }
        upload = {
            "file": (
                "price.xlsx",
                source_bytes,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        }

        denied = client.post(
            "/api/v1/price-imports",
            headers={**viewer_headers, "Idempotency-Key": "browser-key-002"},
            data=form,
            files=upload,
        )
        assert denied.status_code == 403
        created = client.post(
            "/api/v1/price-imports", headers=admin_headers, data=form, files=upload
        )
        assert created.status_code == 202
        import_id = created.json()["import_id"]

        listing = client.get("/api/v1/price-imports", headers=viewer_headers)
        assert listing.status_code == 200
        assert listing.json()["total"] == 1
        summary = listing.json()["items"][0]
        assert summary["id"] == import_id
        assert summary["supplier_external_id"] == "browser-supplier"
        assert summary["filename"] == "price.xlsx"
        original = client.get(
            f"/api/v1/price-versions/{summary['price_version_id']}/original",
            headers=viewer_headers,
        )
        assert original.status_code == 200
        assert original.content == source_bytes
        assert original.headers["content-type"] == (
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        assert original.headers["content-disposition"] == (
            "attachment; filename*=UTF-8''price.xlsx"
        )
        preview = client.get(
            f"/api/v1/price-versions/{summary['price_version_id']}/preview",
            params={"row_offset": 1, "row_limit": 1},
            headers=viewer_headers,
        )
        assert preview.status_code == 200
        assert preview.json() == {
            "version_id": summary["price_version_id"],
            "filename": "price.xlsx",
            "sheets": [
                {"index": 0, "name": "Прайс", "row_count": 3, "column_count": 3},
                {"index": 1, "name": "Условия", "row_count": 1, "column_count": 2},
            ],
            "selected_sheet_index": 0,
            "selected_sheet_name": "Прайс",
            "total_rows": 3,
            "total_columns": 3,
            "row_offset": 1,
            "row_limit": 1,
            "rows": [
                {
                    "row_index": 1,
                    "excel_row": 2,
                    "kind": "source",
                    "cells": [
                        {
                            "column_index": 0,
                            "excel_column": "A",
                            "value": "ФК-1",
                            "is_bold": False,
                            "is_merged_fill": False,
                        },
                        {
                            "column_index": 1,
                            "excel_column": "B",
                            "value": "Фильтр кассетный",
                            "is_bold": False,
                            "is_merged_fill": False,
                        },
                        {
                            "column_index": 2,
                            "excel_column": "C",
                            "value": 1250,
                            "is_bold": False,
                            "is_merged_fill": False,
                        },
                    ],
                }
            ],
        }
        options = client.get("/api/v1/reference/llm-profiles", headers=viewer_headers)
        assert options.status_code == 200
        assert "providers" in options.json()

        # The accounting scope is not an external credential and therefore is
        # absent from the service-client lifecycle API.
        assert client.get("/api/v1/service-clients", headers=admin_headers).json() == []

        async def create_external_client() -> str:
            raw, prefix, digest = generate_service_api_key()
            async with factory() as session:
                session.add(
                    ServiceClient(name="External", key_prefix=prefix, key_digest=digest)
                )
                await session.commit()
            return raw

        external_key = asyncio.run(create_external_client())
        monkeypatch.setattr(settings, "price_import_service_enabled", True)
        assert client.get(
            f"/api/v1/price-versions/{summary['price_version_id']}/original",
            headers={"Authorization": f"Bearer {external_key}"},
        ).status_code == 404
        assert client.get(
            f"/api/v1/price-versions/{summary['price_version_id']}/preview",
            headers={"Authorization": f"Bearer {external_key}"},
        ).status_code == 404
        external_list = client.get(
            "/api/v1/price-imports",
            headers={"Authorization": f"Bearer {external_key}"},
        )
        assert external_list.status_code == 200
        assert external_list.json()["total"] == 0

        async def finish_with_issue() -> None:
            async with factory() as session:
                price_import = (
                    await session.execute(
                        select(PriceImport).where(PriceImport.id == uuid.UUID(import_id))
                    )
                ).scalar_one()
                version = await session.get(PriceVersion, price_import.price_version_id)
                price_import.status = PriceImportStatus.review_required.value
                price_import.stage = PriceImportStatus.review_required.value
                price_import.progress_percent = 100
                version.errors_count = 1
                session.add(
                    PriceParseIssue(
                        price_version_id=version.id,
                        severity="error",
                        code="PRICE_CURRENCY_UNRESOLVED",
                        sheet="Прайс",
                        row_no=2,
                        source_value="100",
                        message="Валюта не определена",
                    )
                )
                await session.commit()

        asyncio.run(finish_with_issue())
        issues = client.get(
            f"/api/v1/price-imports/{import_id}/issues", headers=viewer_headers
        )
        assert issues.status_code == 200
        assert issues.json()["total"] == 1
        assert issues.json()["issues"][0]["code"] == "PRICE_CURRENCY_UNRESOLVED"
    finally:
        app.dependency_overrides.clear()
        asyncio.run(engine.dispose())


def test_import_task_uses_late_ack_for_worker_crash_recovery() -> None:
    from app.tasks import run_price_import

    assert run_price_import.acks_late is True
    assert run_price_import.reject_on_worker_lost is True


def test_deterministic_import_reaches_completed_with_provenance_and_currency(monkeypatch) -> None:
    asyncio.run(_deterministic_import(monkeypatch))


async def _deterministic_import(monkeypatch) -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Прайс"
    sheet.append(["Наименование", "Цена"])
    sheet.append(["Огнетушитель ОП-2", 3400])
    output = io.BytesIO()
    workbook.save(output)
    data = output.getvalue()
    digest = __import__("hashlib").sha256(data).hexdigest()
    monkeypatch.setattr("app.import_service.storage.get_price_file", lambda _key: data)
    engine = create_async_engine("sqlite+aiosqlite://")
    factory = async_sessionmaker(engine, expire_on_commit=False)
    async with engine.begin() as connection:
        await connection.run_sync(Base.metadata.create_all)
    try:
        async with factory() as session:
            client = ServiceClient(name="E2E", key_prefix="e2e", key_digest="e" * 64)
            supplier = Supplier(code="e2e", name="E2E")
            session.add_all([client, supplier])
            await session.flush()
            price_file = PriceFile(
                supplier_id=supplier.id,
                original_filename="price.xlsx",
                storage_key="e2e-file",
                sha256=digest,
                size_bytes=len(data),
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
            session.add(price_file)
            await session.flush()
            version = PriceVersion(
                supplier_id=supplier.id,
                price_file_id=price_file.id,
                version_no=1,
                status="uploaded",
            )
            session.add(version)
            await session.flush()
            price_import = PriceImport(
                service_client_id=client.id,
                supplier_id=supplier.id,
                price_file_id=price_file.id,
                price_version_id=version.id,
                idempotency_key="e2e-import",
                request_fingerprint="e" * 64,
                currency_hint="KZT",
                vat_rate_percent=Decimal("22"),
            )
            session.add(price_import)
            await session.commit()
            status = await advance_price_import(session, price_import.id)
            await session.refresh(price_import)
            values = (await session.execute(select(SupplierPriceValue))).scalars().all()
            items = (await session.execute(select(SupplierPriceItem))).scalars().all()
            assert status == "completed"
            assert price_import.status == "completed"
            assert price_import.completeness_summary["passed"] is True
            assert items[0].source_section_id
            assert items[0].source_panel_index == 0
            assert values[0].currency == "KZT"
            assert values[0].currency_source == "request_hint"
            assert values[0].vat_rate_percent == Decimal("22")
    finally:
        await engine.dispose()


def test_default_currency_resolves_unresolved_prices(monkeypatch) -> None:
    asyncio.run(_default_currency_import(monkeypatch))


async def _default_currency_import(monkeypatch) -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Прайс"
    sheet.append(["Наименование", "Цена"])
    sheet.append(["Огнетушитель ОП-2", 3400])
    output = io.BytesIO()
    workbook.save(output)
    data = output.getvalue()
    digest = __import__("hashlib").sha256(data).hexdigest()
    monkeypatch.setattr("app.import_service.storage.get_price_file", lambda _key: data)
    engine = create_async_engine("sqlite+aiosqlite://")
    factory = async_sessionmaker(engine, expire_on_commit=False)
    async with engine.begin() as connection:
        await connection.run_sync(Base.metadata.create_all)
    try:
        async with factory() as session:
            rub = Currency(code="RUB", name="Российский рубль", symbol="₽", is_default=True)
            client = ServiceClient(name="E2E", key_prefix="e2e", key_digest="e" * 64)
            supplier = Supplier(code="e2e", name="E2E")
            session.add_all([rub, client, supplier])
            await session.flush()
            price_file = PriceFile(
                supplier_id=supplier.id,
                original_filename="price.xlsx",
                storage_key="e2e-file",
                sha256=digest,
                size_bytes=len(data),
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
            session.add(price_file)
            await session.flush()
            version = PriceVersion(
                supplier_id=supplier.id,
                price_file_id=price_file.id,
                version_no=1,
                status="uploaded",
            )
            session.add(version)
            await session.flush()
            price_import = PriceImport(
                service_client_id=client.id,
                supplier_id=supplier.id,
                price_file_id=price_file.id,
                price_version_id=version.id,
                idempotency_key="e2e-import-default",
                request_fingerprint="f" * 64,
                vat_rate_percent=Decimal("22"),
            )
            session.add(price_import)
            await session.commit()
            status = await advance_price_import(session, price_import.id)
            await session.refresh(price_import)
            values = (await session.execute(select(SupplierPriceValue))).scalars().all()
            unresolved = (
                await session.execute(
                    select(PriceParseIssue).where(
                        PriceParseIssue.code == "PRICE_CURRENCY_UNRESOLVED"
                    )
                )
            ).scalars().all()
            assert status == "completed"
            assert price_import.status == "completed"
            assert values[0].currency == "RUB"
            assert values[0].currency_source == "default"
            assert values[0].currency_id == rub.id
            assert values[0].vat_rate_percent == Decimal("22")
            assert unresolved == []
    finally:
        await engine.dispose()


def test_quality_gate_blocks_without_vat_rate(tmp_path, monkeypatch) -> None:
    """Import cannot complete without vat_rate_percent set."""

    async def _test():
        from app.config import settings
        from app.import_service import run_quality_gate

        database_path = tmp_path / "vat_block.sqlite"
        engine = create_async_engine(f"sqlite+aiosqlite:///{database_path}")
        factory = async_sessionmaker(engine, expire_on_commit=False)
        monkeypatch.setattr(settings, "price_import_service_enabled", True)

        try:
            async with engine.begin() as connection:
                await connection.run_sync(Base.metadata.create_all)

            async with factory() as session:
                client = ServiceClient(
                    name="test client",
                    key_prefix="test",
                    key_digest="digest",
                )
                session.add(client)
                await session.flush()

                supplier = Supplier(
                    service_client_id=client.id,
                    external_id="test-supplier",
                    code="test-supplier",
                    name="Test Supplier",
                )
                session.add(supplier)
                await session.flush()

                price_file = PriceFile(
                    project_id=uuid.UUID("10000000-0000-4000-8000-000000000001"),
                    supplier_id=supplier.id,
                    storage_key="test/key",
                    sha256="a" * 64,
                    original_filename="test.xlsx",
                    size_bytes=100,
                    mime="application/vnd.ms-excel",
                )
                session.add(price_file)
                await session.flush()

                version = PriceVersion(
                    supplier_id=supplier.id,
                    price_file_id=price_file.id,
                    version_no=1,
                    status=VersionStatus.validated.value,
                    loader="xlsx",
                )
                session.add(version)
                await session.flush()

                # НДС не указан (None)
                price_import = PriceImport(
                    service_client_id=client.id,
                    supplier_id=supplier.id,
                    price_file_id=price_file.id,
                    price_version_id=version.id,
                    idempotency_key="no-vat-test",
                    request_fingerprint="b" * 64,
                    vat_rate_percent=None,  # НДС не указан
                )
                session.add(price_import)
                await session.commit()

                # Попытка пройти quality gate без НДС
                status = await run_quality_gate(
                    session,
                    price_import=price_import,
                    version=version,
                    price_file=price_file,
                    data=b"fake data",
                )

                await session.refresh(price_import)

                # Проверка: импорт заблокирован
                assert status == "review_required"
                assert price_import.status == "review_required"
                assert price_import.error_code == "VAT_RATE_REQUIRED"
                assert "ставка НДС".lower() in price_import.error_message.lower()
        finally:
            await engine.dispose()

    asyncio.run(_test())


def test_patch_vat_rate_updates_import(tmp_path, monkeypatch) -> None:
    """PATCH /price-imports/{id} with vat_rate_percent updates the import."""

    from app.config import settings
    from app.main import app
    import app.api.price_imports as api_module
    import app.tasks as tasks_module

    database_path = tmp_path / "patch_vat.sqlite"
    engine = create_async_engine(f"sqlite+aiosqlite:///{database_path}")
    factory = async_sessionmaker(engine, expire_on_commit=False)
    monkeypatch.setattr(settings, "price_import_service_enabled", True)
    monkeypatch.setattr(settings, "service_api_key_pepper", "p" * 48)

    class FakeTask:
        @staticmethod
        def delay(*_args, **_kwargs):
            return None

    monkeypatch.setattr(tasks_module, "advance_price_import_task", FakeTask())

    async def session_override():
        async with factory() as session:
            yield session

    async def prepare() -> tuple[str, uuid.UUID]:
        async with engine.begin() as connection:
            await connection.run_sync(Base.metadata.create_all)
        raw, prefix, digest = generate_service_api_key()
        async with factory() as session:
            client = ServiceClient(name="patch client", key_prefix=prefix, key_digest=digest)
            session.add(client)
            await session.flush()

            supplier = Supplier(
                service_client_id=client.id,
                external_id="patch-supplier",
                code="patch-supplier",
                name="Patch Supplier",
            )
            session.add(supplier)
            await session.flush()

            price_file = PriceFile(
                project_id=uuid.UUID("10000000-0000-4000-8000-000000000001"),
                supplier_id=supplier.id,
                storage_key="patch/key",
                sha256="c" * 64,
                original_filename="patch.xlsx",
                size_bytes=200,
                mime="application/vnd.ms-excel",
            )
            session.add(price_file)
            await session.flush()

            version = PriceVersion(
                supplier_id=supplier.id,
                price_file_id=price_file.id,
                version_no=1,
                status=VersionStatus.validated.value,
                loader="xlsx",
            )
            session.add(version)
            await session.flush()

            price_import = PriceImport(
                service_client_id=client.id,
                supplier_id=supplier.id,
                price_file_id=price_file.id,
                price_version_id=version.id,
                idempotency_key="patch-test",
                request_fingerprint="d" * 64,
                status=PriceImportStatus.review_required.value,
                error_code="VAT_RATE_REQUIRED",
                vat_rate_percent=None,
            )
            session.add(price_import)
            await session.commit()
            return raw, price_import.id

    raw_key, import_id = asyncio.run(prepare())
    app.dependency_overrides[get_session] = session_override
    try:
        client = TestClient(app)
        headers = {"Authorization": f"Bearer {raw_key}"}

        # PATCH с vat_rate_percent
        response = client.patch(
            f"/api/v1/price-imports/{import_id}",
            headers=headers,
            data={"vat_rate_percent": "22"},
        )

        assert response.status_code == 200
        result = response.json()

        # Проверка: НДС установлен
        async def check_vat():
            async with factory() as session:
                price_import = await session.get(PriceImport, import_id)
                return price_import.vat_rate_percent

        vat = asyncio.run(check_vat())
        assert vat == Decimal("22")
    finally:
        app.dependency_overrides.clear()
        asyncio.run(engine.dispose())


def test_patch_vat_rate_backfills_existing_supplier_price_values(

    tmp_path, monkeypatch
) -> None:
    """P0-3: PATCHing the VAT rate backfills already-persisted values.

    SupplierPriceValue rows are materialised before the operator can choose
    the import rate (the VAT_RATE_REQUIRED gate), so a PATCH must atomically
    fill SupplierPriceValue.vat_rate_percent for the values of the import
    version that do not carry a rate yet.  Explicit value-level rates are
    preserved, values of other versions are untouched and a repeated PATCH
    is idempotent.
    """

    from app.config import settings
    from app.main import app
    import app.tasks as tasks_module
    from app.models import VatMode

    database_path = tmp_path / "patch_vat_backfill.sqlite"
    engine = create_async_engine(f"sqlite+aiosqlite:///{database_path}")
    factory = async_sessionmaker(engine, expire_on_commit=False)
    monkeypatch.setattr(settings, "price_import_service_enabled", True)
    monkeypatch.setattr(settings, "service_api_key_pepper", "p" * 48)

    class FakeTask:
        @staticmethod
        def delay(*_args, **_kwargs):
            return None

    monkeypatch.setattr(tasks_module, "advance_price_import_task", FakeTask())

    async def session_override():
        async with factory() as session:
            yield session

    async def prepare() -> tuple[str, uuid.UUID]:
        async with engine.begin() as connection:
            await connection.run_sync(Base.metadata.create_all)
        raw, prefix, digest = generate_service_api_key()
        async with factory() as session:
            client = ServiceClient(name="backfill client", key_prefix=prefix, key_digest=digest)
            session.add(client)
            await session.flush()

            supplier = Supplier(
                service_client_id=client.id,
                external_id="backfill-supplier",
                code="backfill-supplier",
                name="Backfill Supplier",
            )
            session.add(supplier)
            await session.flush()

            price_file = PriceFile(
                project_id=uuid.UUID("10000000-0000-4000-8000-000000000001"),
                supplier_id=supplier.id,
                storage_key="backfill/key",
                sha256="f" * 64,
                original_filename="backfill.xlsx",
                size_bytes=200,
                mime="application/vnd.ms-excel",
            )
            session.add(price_file)
            await session.flush()

            version = PriceVersion(
                supplier_id=supplier.id,
                price_file_id=price_file.id,
                version_no=1,
                status=VersionStatus.validated.value,
                loader="xlsx",
            )
            session.add(version)
            await session.flush()

            other_version = PriceVersion(
                supplier_id=supplier.id,
                price_file_id=price_file.id,
                version_no=2,
                status=VersionStatus.validated.value,
                loader="xlsx",
            )
            session.add(other_version)
            await session.flush()

            price_import = PriceImport(
                service_client_id=client.id,
                supplier_id=supplier.id,
                price_file_id=price_file.id,
                price_version_id=version.id,
                idempotency_key="patch-backfill-test",
                request_fingerprint="e" * 64,
                status=PriceImportStatus.review_required.value,
                error_code="VAT_RATE_REQUIRED",
                vat_rate_percent=None,
            )
            session.add(price_import)
            await session.flush()

            # Import-version positions created before the operator picked the
            # rate (VAT_RATE_REQUIRED): the rate is still missing on them.
            missing_item = SupplierPriceItem(
                price_version_id=version.id,
                source_name="missing",
                normalized_name="missing",
                source_sheet="sheet",
                source_row=2,
            )
            session.add(missing_item)
            await session.flush()
            for price_kind in ("wholesale", "retail"):
                session.add(
                    SupplierPriceValue(
                        item=missing_item,
                        price_kind=price_kind,
                        amount=Decimal("100"),
                        vat_mode=VatMode.unknown.value,
                        vat_rate_percent=None,
                    )
                )

            # An explicit value-level rate coming from the parser must not be
            # overwritten by the import-level backfill.
            explicit_item = SupplierPriceItem(
                price_version_id=version.id,
                source_name="explicit",
                normalized_name="explicit",
                source_sheet="sheet",
                source_row=3,
            )
            session.add(explicit_item)
            await session.flush()
            session.add(
                SupplierPriceValue(
                    item=explicit_item,
                    price_kind="wholesale",
                    amount=Decimal("200"),
                    vat_mode=VatMode.included.value,
                    vat_rate_percent=Decimal("10"),
                )
            )

            # A value of another version of the same supplier must stay as is.
            other_item = SupplierPriceItem(
                price_version_id=other_version.id,
                source_name="other",
                normalized_name="other",
                source_sheet="sheet",
                source_row=4,
            )
            session.add(other_item)
            await session.flush()
            session.add(
                SupplierPriceValue(
                    item=other_item,
                    price_kind="wholesale",
                    amount=Decimal("300"),
                    vat_mode=VatMode.unknown.value,
                    vat_rate_percent=None,
                )
            )
            await session.commit()
            return raw, price_import.id

    raw_key, import_id = asyncio.run(prepare())
    app.dependency_overrides[get_session] = session_override
    try:
        client = TestClient(app)
        headers = {"Authorization": f"Bearer {raw_key}"}

        async def value_rates() -> dict[str, Decimal | None]:
            async with factory() as session:
                rows = (
                    await session.execute(
                        select(
                            SupplierPriceItem.source_name,
                            SupplierPriceValue.price_kind,
                            SupplierPriceValue.vat_rate_percent,
                        )
                        .join(
                            SupplierPriceValue,
                            SupplierPriceValue.item_id == SupplierPriceItem.id,
                        )
                        .order_by(SupplierPriceItem.source_row, SupplierPriceValue.price_kind)
                    )
                ).all()
            return {f"{name}:{kind}": rate for name, kind, rate in rows}

        assert asyncio.run(value_rates()) == {
            "missing:retail": None,
            "missing:wholesale": None,
            "explicit:wholesale": Decimal("10"),
            "other:wholesale": None,
        }

        # PATCH 22: the rate must reach already-persisted values.
        response = client.patch(
            f"/api/v1/price-imports/{import_id}",
            headers=headers,
            data={"vat_rate_percent": "22"},
        )
        assert response.status_code == 200, response.text

        assert asyncio.run(value_rates()) == {
            "missing:retail": Decimal("22"),
            "missing:wholesale": Decimal("22"),
            "explicit:wholesale": Decimal("10"),
            "other:wholesale": None,
        }

        # A repeated PATCH is idempotent: values keep their rate.
        response = client.patch(
            f"/api/v1/price-imports/{import_id}",
            headers=headers,
            data={"vat_rate_percent": "22"},
        )
        assert response.status_code == 200, response.text
        assert asyncio.run(value_rates()) == {
            "missing:retail": Decimal("22"),
            "missing:wholesale": Decimal("22"),
            "explicit:wholesale": Decimal("10"),
            "other:wholesale": None,
        }

        async def stored_rate() -> Decimal | None:
            async with factory() as session:
                price_import = await session.get(PriceImport, import_id)
                return price_import.vat_rate_percent

        assert asyncio.run(stored_rate()) == Decimal("22")
    finally:
        app.dependency_overrides.clear()
        asyncio.run(engine.dispose())


def test_vat_rate_patch_resume_runs_only_quality_gate(monkeypatch) -> None:
    """P0-3 resume: PATCH(rate) + advance on a validated deterministic import
    must re-run the quality gates, not start LLM rule-generation attempts."""

    async def scenario() -> None:
        import app.import_service as import_module
        from app.import_service import backfill_vat_rate_percent

        async def money_gate(*_args, **_kwargs):
            return {"passed": True, "failures": 0}

        async def completeness_gate(*_args, **_kwargs):
            return {"passed": True, "confirmed_missed_positions": 0}

        async def category_gate(*_args, **_kwargs):
            return {"passed": True, "failures": 0}

        async def telemetry(*_args, **_kwargs):
            return {}

        monkeypatch.setattr(import_module, "apply_money_gate", money_gate)
        monkeypatch.setattr(import_module, "reconcile_version", completeness_gate)
        monkeypatch.setattr(import_module, "apply_category_context_gate", category_gate)
        monkeypatch.setattr(import_module, "aggregate_import_telemetry", telemetry)
        monkeypatch.setattr("app.import_service.storage.get_price_file", lambda _key: b"")

        engine = create_async_engine("sqlite+aiosqlite://")
        factory = async_sessionmaker(engine, expire_on_commit=False)
        async with engine.begin() as connection:
            await connection.run_sync(Base.metadata.create_all)
        try:
            async with factory() as session:
                client = ServiceClient(
                    name="vat resume", key_prefix="vatresume", key_digest="r" * 64
                )
                supplier = Supplier(code="vat-resume", name="VAT Resume")
                session.add_all([client, supplier])
                await session.flush()
                price_file = PriceFile(
                    supplier_id=supplier.id,
                    original_filename="vat-resume.xlsx",
                    storage_key="vat-resume-file",
                    sha256="v" * 64,
                    size_bytes=1,
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
                session.add(price_file)
                await session.flush()
                version = PriceVersion(
                    supplier_id=supplier.id,
                    price_file_id=price_file.id,
                    version_no=1,
                    status=VersionStatus.validated.value,
                    item_count=1,
                    rows_accepted=1,
                    errors_count=0,
                    warnings_count=0,
                )
                session.add(version)
                await session.flush()
                price_import = PriceImport(
                    service_client_id=client.id,
                    supplier_id=supplier.id,
                    price_file_id=price_file.id,
                    price_version_id=version.id,
                    idempotency_key="vat-resume-import",
                    request_fingerprint="v" * 64,
                    status=PriceImportStatus.queued.value,
                    stage=PriceImportStatus.queued.value,
                    vat_rate_percent=None,
                )
                session.add(price_import)
                await session.flush()
                item = SupplierPriceItem(
                    price_version_id=version.id,
                    source_name="Клапан",
                    normalized_name="Клапан",
                    source_sheet="Прайс",
                    source_row=2,
                )
                session.add(item)
                await session.flush()
                session.add(
                    SupplierPriceValue(
                        item=item,
                        price_kind="wholesale",
                        amount=Decimal("618.40"),
                        vat_mode="unknown",
                        vat_rate_percent=None,
                    )
                )
                await session.commit()

                # Run 1: no rate -> the quality gate blocks with VAT_RATE_REQUIRED.
                first = await advance_price_import(session, price_import.id)
                await session.refresh(price_import)
                assert first == "review_required"
                assert price_import.status == "review_required"
                assert price_import.error_code == "VAT_RATE_REQUIRED"
                attempts_after_first = (
                    await session.execute(
                        select(func.count()).select_from(RuleGenerationAttempt)
                    )
                ).scalar_one()
                assert attempts_after_first == 0

                # Emulate the PATCH handler: record the rate, backfill the
                # materialised values and move the import back to quality_check.
                price_import.vat_rate_percent = Decimal("22")
                price_import.status = PriceImportStatus.quality_check.value
                price_import.error_code = None
                price_import.error_message = None
                await backfill_vat_rate_percent(
                    session,
                    price_version_id=version.id,
                    vat_rate_percent=Decimal("22"),
                )
                await session.commit()

                # Run 2: only the quality gates must run; no LLM attempts, no
                # rule generation, and the deterministic import completes.
                second = await advance_price_import(session, price_import.id)
                await session.refresh(price_import)
                assert second == "completed", (
                    f"second={second} error_code={price_import.error_code} "
                    f"error_message={price_import.error_message}"
                )
                assert price_import.status == "completed"
                attempts = (
                    await session.execute(
                        select(func.count()).select_from(RuleGenerationAttempt)
                    )
                ).scalar_one()
                assert attempts == 0
                value = (
                    await session.execute(select(SupplierPriceValue))
                ).scalar_one()
                assert value.vat_rate_percent == Decimal("22")
        finally:
            await engine.dispose()

    asyncio.run(scenario())
