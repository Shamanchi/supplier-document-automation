from __future__ import annotations

from decimal import Decimal
from io import BytesIO

from openpyxl import load_workbook

from app.parsing.pipeline import ParsedAttribute, ParsedItem, ParsedValue
from app.parsing.validation_export import build_validation_export, build_validation_xlsx


def _item(*, row: int, category: str, name: str) -> ParsedItem:
    return ParsedItem(
        source_name=name,
        normalized_name=name.casefold(),
        source_article="A-1",
        normalized_article="a-1",
        unit_code="шт",
        category_path=category,
        dimensions=None,
        description=None,
        attribute_values=[
            ParsedAttribute(
                source_header="Тип шкафа", raw_value="закрытый", role="attribute",
                attribute_code=None, value_kind="text", canonical_unit=None,
                ordinal=0, source_column_index=1,
            )
        ],
        source_sheet="Шкафы",
        source_row=row,
        source_section_id="tpl_shkafy",
        source_header_cells=[[0, "Наименование"], [1, "Тип шкафа"], [2, "Цена"]],
        values=[
            ParsedValue(
                kind="wholesale", amount=Decimal("123.45"), vat_mode="included",
                is_on_request=False, source_label="Цена", currency="RUB",
                source_column_index=2,
            )
        ],
    )


def test_json_export_contains_every_position_grouped_by_category() -> None:
    payload = build_validation_export([
        _item(row=7, category="Компакт", name="ШПК-1"),
        _item(row=21, category="Универсальные", name="ШПК-2"),
    ], filename="price.xlsx")

    assert payload["filename"] == "price.xlsx"
    assert payload["item_count"] == 2
    assert [section["category"] for section in payload["sections"]] == [
        "Компакт", "Универсальные",
    ]
    assert payload["sections"][0]["source_header"] == [
        {"column_index": 0, "value": "Наименование"},
        {"column_index": 1, "value": "Тип шкафа"},
        {"column_index": 2, "value": "Цена"},
    ]


def test_xlsx_export_omits_source_headers_and_autofits_content() -> None:
    content = build_validation_xlsx([
        _item(row=7, category="Компакт", name="ШПК-1"),
        _item(row=21, category="Универсальные", name="ШПК-2"),
    ], filename="price.xlsx")

    workbook = load_workbook(BytesIO(content), data_only=True)
    assert workbook.sheetnames == ["Результат разбора"]
    rows = list(workbook.active.iter_rows(values_only=True))
    assert rows[0][0] == "Категория"
    assert rows[0][1] == "Компакт"
    assert rows[1][0] == "Лист"
    assert rows[2][3] == "ШПК-1"
    assert not any(row[0] == "Исходный заголовок" for row in rows)
    assert any(all(value is None for value in row) for row in rows[3:6])
    assert any(row[0] == "Категория" and row[1] == "Универсальные" for row in rows)
    assert workbook.active.column_dimensions["D"].width < 60
    assert workbook.active.column_dimensions["G"].width > workbook.active.column_dimensions["B"].width


def test_xlsx_export_preserves_qualified_price_meaning() -> None:
    ranged = _item(row=7, category="Компакт", name="Диапазон")
    ranged.values = [
        ParsedValue(
            kind="wholesale", amount=Decimal("130.00"), qualifier="range",
            amount_from=Decimal("130.00"), amount_to=Decimal("300.00"),
            vat_mode="included", is_on_request=False, source_label="Цена",
            currency="RUB", raw_value="130-300", source_column_index=2,
        )
    ]
    lower_bound = _item(row=8, category="Компакт", name="Нижняя граница")
    lower_bound.values = [
        ParsedValue(
            kind="wholesale", amount=Decimal("1000.00"), qualifier="from",
            amount_from=Decimal("1000.00"), vat_mode="included",
            is_on_request=False, source_label="Цена", currency="RUB",
            raw_value="от 1 000 р", source_column_index=2,
        )
    ]
    rate = _item(row=9, category="Компакт", name="Тариф")
    rate.values = [
        ParsedValue(
            kind="wholesale", amount=Decimal("50.00"), amount_from=Decimal("50.00"),
            vat_mode="included", is_on_request=False, source_label="Цена",
            currency="KZT", basis_quantity=Decimal("1.00"), basis_unit="км",
            raw_value="50 тенге за 1 км", source_column_index=2,
        )
    ]

    content = build_validation_xlsx(
        [ranged, lower_bound, rate], filename="qualified.xlsx"
    )
    workbook = load_workbook(BytesIO(content), data_only=True)
    prices_by_name = {
        row[3]: row[7]
        for row in workbook.active.iter_rows(values_only=True)
        if row[3] in {"Диапазон", "Нижняя граница", "Тариф"}
    }

    assert prices_by_name == {
        "Диапазон": "Цена: 130.00–300.00 RUB",
        "Нижняя граница": "Цена: от 1000.00 RUB",
        "Тариф": "Цена: 50.00 KZT за 1.00 км",
    }
