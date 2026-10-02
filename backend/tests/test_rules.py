"""Parse-rule mechanics: fingerprint stability, detection, recipe roundtrip,
auto-draft creation after a successful parse (sqlite, no MinIO)."""
from __future__ import annotations

import asyncio
import io
import json
from decimal import Decimal
from pathlib import Path

import pytest
from fastapi import HTTPException
from fastapi.testclient import TestClient
from sqlalchemy import select
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

from app.config import settings
from app.db import get_session
from app.main import app
from app import storage
from app.api.versions import _export_json, list_items
from app.api.parse_rules import (
    _validate_recipe_format,
    set_rule_status,
    auto_draft_recipe,
    clone_rule_to_draft,
    preview_rule_matrix,
    rebind_rule_matrix,
    update_rule,
    validate_rule,
)
from app.models import (
    ARCHIVE_PROJECT_ID,
    ORIGINAL_PROJECT_ID,
    Base,
    GlobalCategory,
    ParseRule,
    ParseRuleProvenance,
    ParseRuleStatus,
    RuleMatch,
    PriceFile,
    PriceVersion,
    Project,
    Supplier,
    SupplierCategory,
    SupplierCategoryMapping,
    SupplierPriceItem,
    VersionStatus,
)
from app.parse_service import parse_and_persist
from app.parsing.grid import Cell, SheetGrid
from app.parsing.mapping import ColumnMapping, MappingDecision, TableMapping, build_auto_mapping
from app.parsing.pipeline import parse_grids
from app.parsing.rules import (
    compute_fingerprint,
    compute_header_hash,
    jaccard,
    partial_similarity,
    recipe_from_decision,
    recipe_mapping_gaps,
    recipe_prefill_mapping,
    recipe_to_decision,
)
from app.schemas import (
    ParseRuleStatusUpdate,
    ParseRuleUpdate,
    RuleMatrixPreviewIn,
    RuleMatrixRebindIn,
    RuleValidateIn,
)


def test_renaming_rule_preserves_successful_dry_run() -> None:
    asyncio.run(_rename_rule_preserves_successful_dry_run())


async def _rename_rule_preserves_successful_dry_run() -> None:
    engine = create_async_engine("sqlite+aiosqlite://")
    factory = async_sessionmaker(engine, expire_on_commit=False)
    async with engine.begin() as connection:
        await connection.run_sync(Base.metadata.create_all)
    try:
        async with factory() as session:
            project = Project(id=ORIGINAL_PROJECT_ID, slug="original-prices", name="Оригинальные прайсы", is_default=True)
            supplier = Supplier(code="rename-rule", name="Rename rule")
            session.add_all([project, supplier])
            await session.flush()
            validation = {"passed": True, "metrics": {"items_with_price": 10}}
            rule = ParseRule(
                project_id=project.id, supplier_id=supplier.id, name="Первый лист",
                loader="xlsx", header_hash="d" * 64, fingerprint={}, recipe={"sheets": []},
                revision=1, status="draft", provenance="model_proposed", validation=validation,
            )
            session.add(rule)
            await session.flush()

            result = await update_rule(
                rule.id,
                ParseRuleUpdate(name="price_list_sample"),
                project=project,
                session=session,
            )

            assert result.name == "price_list_sample"
            assert result.validation == validation
            assert result.revision == 2
    finally:
        await engine.dispose()


def _grid(headers: list[str], data_rows: list[list[object]], name: str = "Лист1") -> SheetGrid:
    rows = [[Cell(value=h) for h in headers]]
    for raw in data_rows:
        rows.append([Cell(value=v) for v in raw])
    return SheetGrid(name=name, rows=rows)


HEADERS = ["Артикул", "Наименование", "Ед. изм.", "Цена"]
ROWS_A = [["A-1", "Муфта 1", "шт", 100.5], ["A-2", "Муфта 2", "шт", 200]]
ROWS_B = [["B-9", "Клапан", "шт", 33], ["B-10", "Кран", "шт", 44], ["B-11", "Затвор", "шт", 55]]
COMPACT_DIMENSIONS_HEADERS = ["Наименование", "Шир*Выс*Гл", "Цена"]
COMPACT_DIMENSIONS_ROWS = [["ШПК-310Н", "540х640х230", 1188]]
TARGET_HEADERS = [
    "Наименование",
    "Артикул",
    "Толщина металла",
    "Габаритные размеры Длина/глубина/высота, мм",
    "Описание",
    "Цена, с НДС",
]
TARGET_ROWS = [[
    "Пост пожарный",
    "ППМП-02",
    "Лист 0.8 мм",
    "1000х900х1670",
    "Передвижной пост для средств пожаротушения",
    "64 500",
]]

PRICE_NEW_FILE = (
    Path(__file__).parents[4]
    / "original"
    / "datasets"
    / "прайсы"
    / "price_list_sample.xlsx"
)
requires_price_new = pytest.mark.skipif(
    not PRICE_NEW_FILE.is_file(), reason="price_new corpus file is not available"
)


def test_parse_rule_page_filters_and_paginates(tmp_path, monkeypatch) -> None:
    database_path = tmp_path / "rule-page.sqlite"
    engine = create_async_engine(f"sqlite+aiosqlite:///{database_path}")
    factory = async_sessionmaker(engine, expire_on_commit=False)
    ids: dict[str, str] = {}
    for key, value in {
        "app_admin_username": "rule-admin",
        "app_admin_password": "rule-admin-password-long",
        "app_user_username": "rule-viewer",
        "app_user_password": "rule-viewer-password-long",
        "app_auth_secret": "rule-auth-secret-that-is-long-enough-12345",
        "app_auth_secure_cookies": False,
    }.items():
        monkeypatch.setattr(settings, key, value)

    async def session_override():
        async with factory() as session:
            yield session

    async def prepare() -> None:
        async with engine.begin() as connection:
            await connection.run_sync(Base.metadata.create_all)
        async with factory() as session:
            session.add_all(
                [
                    Project(
                        id=ORIGINAL_PROJECT_ID,
                        slug="original-prices",
                        name="Оригинальные прайсы",
                        is_default=True,
                    ),
                    Project(
                        id=ARCHIVE_PROJECT_ID,
                        slug="legacy-archive",
                        name="Архив",
                        is_archived=True,
                    ),
                ]
            )
            supplier = Supplier(code="rule-page", name="Поставщик правил")
            session.add(supplier)
            await session.flush()
            active_rule = ParseRule(
                supplier_id=supplier.id,
                name="Правило XLSX",
                loader="xlsx",
                header_hash="a" * 64,
                fingerprint={},
                recipe={"sheets": []},
                revision=1,
                status="active",
                provenance="manual",
            )
            session.add_all(
                [active_rule, ParseRule(
                    supplier_id=supplier.id,
                    name="Черновик",
                    loader="xls",
                    header_hash="b" * 64,
                    fingerprint={},
                    recipe={"sheets": []},
                    revision=1,
                    status="draft",
                    provenance="manual",
                ), ParseRule(
                    project_id=ARCHIVE_PROJECT_ID,
                    supplier_id=supplier.id,
                    name="Архивное правило",
                    loader="xlsx",
                    header_hash="c" * 64,
                    fingerprint={},
                    recipe={"sheets": []},
                    revision=1,
                    status="active",
                    provenance="manual",
                )]
            )
            await session.flush()
            archive_rule = (
                await session.execute(
                    select(ParseRule).where(ParseRule.project_id == ARCHIVE_PROJECT_ID)
                )
            ).scalar_one()
            ids["archive_rule"] = str(archive_rule.id)
            ids["active_rule"] = str(active_rule.id)
            price_file = PriceFile(
                supplier_id=supplier.id,
                original_filename="rule-price.xlsx",
                storage_key="prices/rule-price.xlsx",
                sha256="c" * 64,
                size_bytes=100,
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
            session.add(price_file)
            await session.flush()
            version = PriceVersion(
                supplier_id=supplier.id,
                price_file_id=price_file.id,
                version_no=1,
                status="validated",
                parse_rule_id=active_rule.id,
                item_count=3,
                rows_accepted=3,
            )
            session.add(version)
            await session.flush()
            active_rule.created_from_version_id = version.id
            session.add_all([
                SupplierPriceItem(
                    price_version_id=version.id,
                    source_name=f"Позиция {index + 1}",
                    normalized_name=f"позиция {index + 1}",
                    source_sheet="Прайс",
                    source_row=source_row,
                    source_panel_index=panel_index,
                )
                for index, (source_row, panel_index) in enumerate([(2, 0), (2, 1), (3, None)])
            ])
            await session.commit()

    asyncio.run(prepare())
    app.dependency_overrides[get_session] = session_override
    try:
        client = TestClient(app)
        token = client.post(
            "/api/v1/auth/login",
            json={"username": "rule-viewer", "password": "rule-viewer-password-long"},
        ).json()["access_token"]
        response = client.get(
            "/api/v1/parse-rules",
            headers={"Authorization": f"Bearer {token}"},
            params={"status": "active", "query": "Поставщик", "limit": 1},
        )
        assert response.status_code == 200
        assert response.json()["total"] == 1
        assert response.json()["items"][0]["name"] == "Правило XLSX"
        assert response.json()["items"][0]["supplier_name"] == "Поставщик правил"
        assert response.json()["items"][0]["usage"] == {
            "price_versions": 1,
            "database_items": 3,
            "source_rows": 2,
        }
        archive_response = client.get(
            "/api/v1/parse-rules",
            headers={
                "Authorization": f"Bearer {token}",
                "X-Project-ID": str(ARCHIVE_PROJECT_ID),
            },
        )
        assert archive_response.status_code == 200
        assert archive_response.json()["total"] == 1
        assert archive_response.json()["items"][0]["name"] == "Архивное правило"
        detail = client.get(
            f"/api/v1/parse-rules/{ids['active_rule']}",
            headers={"Authorization": f"Bearer {token}"},
        )
        assert detail.status_code == 200
        detail_body = detail.json()
        assert detail_body["name"] == "Правило XLSX"
        assert detail_body["supplier_name"] == "Поставщик правил"
        assert detail_body["usage"] == {
            "price_versions": 1,
            "database_items": 3,
            "source_rows": 2,
        }
        assert detail_body["source_version"] == {
            "id": str(detail_body["created_from_version_id"]),
            "supplier_id": str(response.json()["items"][0]["supplier_id"]),
            "supplier_external_id": None,
            "supplier_name": "Поставщик правил",
            "version_no": 1,
            "status": "validated",
            "loader": None,
            "filename": "rule-price.xlsx",
            "item_count": 3,
            "warnings_count": 0,
            "errors_count": 0,
            "created_at": detail_body["source_version"]["created_at"],
            "import_id": None,
        }
        hidden_detail = client.get(
            f"/api/v1/parse-rules/{ids['archive_rule']}",
            headers={"Authorization": f"Bearer {token}"},
        )
        assert hidden_detail.status_code == 404
    finally:
        app.dependency_overrides.clear()
        asyncio.run(engine.dispose())


def test_fingerprint_stable_across_different_rows():
    fp_a = compute_fingerprint("xlsx", [_grid(HEADERS, ROWS_A)])
    fp_b = compute_fingerprint("xlsx", [_grid(HEADERS, ROWS_B)])
    assert fp_a == fp_b
    assert compute_header_hash(fp_a) == compute_header_hash(fp_b)


def test_fingerprint_changes_with_headers_and_loader():
    fp = compute_fingerprint("xlsx", [_grid(HEADERS, ROWS_A)])
    fp_other = compute_fingerprint(
        "xlsx", [_grid(["Артикул", "Наименование", "Ед. изм.", "Цена опт"], ROWS_A)]
    )
    assert compute_header_hash(fp) != compute_header_hash(fp_other)

    fp_pdf = compute_fingerprint("pdf_text", [_grid(HEADERS, ROWS_A)])
    assert compute_header_hash(fp) != compute_header_hash(fp_pdf)


def test_fingerprint_includes_each_distinct_vertical_table_header():
    first_section = [
        [Cell(value="Наименование"), Cell(value="Цена")],
        [Cell(value="Шкаф"), Cell(value="100")],
    ]
    full = SheetGrid(
        name="Прайс",
        rows=first_section
        + [
            [Cell(value=None), Cell(value=None)],
            [Cell(value="Кол-во секций"), Cell(value="Опт-1")],
            [Cell(value="2"), Cell(value="200")],
        ],
    )
    short = SheetGrid(name="Прайс", rows=first_section)

    full_fingerprint = compute_fingerprint("xlsx", [full])
    short_fingerprint = compute_fingerprint("xlsx", [short])

    assert full_fingerprint != short_fingerprint
    assert compute_header_hash(full_fingerprint) != compute_header_hash(short_fingerprint)

    # Defence in depth for recipes created before the fuller fingerprint: the
    # missing second section is skipped rather than indexing a non-existent
    # source row.
    recipe = recipe_from_decision(build_auto_mapping([full]), [full])
    result = parse_grids([short], recipe_to_decision(recipe, [short]))
    assert [item.source_name for item in result.items] == ["Шкаф"]


def test_partial_similarity_thresholds():
    fp = compute_fingerprint("xlsx", [_grid(HEADERS, ROWS_A)])
    # one header of four renamed -> jaccard 3/5 = 0.6 -> partial
    fp_sim = compute_fingerprint(
        "xlsx", [_grid(["Артикул", "Наименование", "Ед. изм.", "Розничная цена"], ROWS_A)]
    )
    score = partial_similarity(fp, fp_sim)
    assert score == pytest.approx(0.6)
    fp_diff = compute_fingerprint("xlsx", [_grid(["Фото", "Описание", "Вес", "Склад"], ROWS_A)])
    # completely different headers, but detect_header needs a name column
    assert partial_similarity(fp, fp_diff) < 0.6
    assert jaccard(set("ab"), set("ab")) == 1.0


def test_recipe_mapping_roundtrip():
    grid = _grid(HEADERS, ROWS_A, name="Крепёж")
    decision = build_auto_mapping([grid])
    assert decision.auto_ok

    recipe = recipe_from_decision(decision, [grid])
    assert recipe["options"]["currency"] is None
    assert recipe["sheets"][0]["header_row"] == 0
    assert recipe["sheets"][0]["default_category_source"] == "sheet_name"
    assert recipe["sheets"][0]["default_category_path"] is None

    decision2 = recipe_to_decision(recipe, [grid])
    assert decision2.auto_ok
    orig = {(c.index, c.role) for c in decision.tables[0].columns}
    restored = {(c.index, c.role) for c in decision2.tables[0].columns}
    assert orig == restored
    assert decision2.tables[0].default_category_path == "Крепёж"
    # price descriptor survives
    price_col = [c for c in decision2.tables[0].columns if c.role == "price"][0]
    assert price_col.price_descriptor is not None


def test_recipe_omits_blank_trailing_ignore_columns_but_keeps_internal_spacers():
    grid = _grid(
        ["Наименование", "Цена", "", "Наименование", "Цена", ""],
        [["Товар А", 10, None, "Товар Б", 20, None]],
    )
    decision = MappingDecision(
        auto_ok=True,
        tables=[
            TableMapping(
                sheet_name=grid.name,
                header_row_index=0,
                columns=[
                    ColumnMapping(index=0, role="name", header="Наименование"),
                    ColumnMapping(index=1, role="price", header="Цена"),
                    ColumnMapping(index=2, role="ignore", header=""),
                    ColumnMapping(index=3, role="name", header="Наименование"),
                    ColumnMapping(index=4, role="price", header="Цена"),
                    ColumnMapping(index=5, role="ignore", header=""),
                ],
            )
        ],
    )

    recipe = recipe_from_decision(decision, [grid])

    assert [column["index"] for column in recipe["sheets"][0]["columns"]] == [
        0,
        1,
        2,
        3,
        4,
    ]


def test_recipe_sheet_category_source_follows_the_resolved_sheet_name():
    original = _grid(HEADERS, ROWS_A, name="Крепёж")
    recipe = recipe_from_decision(build_auto_mapping([original]), [original])
    renamed = _grid(HEADERS, ROWS_A, name="Крепёж новый")

    restored = recipe_to_decision(recipe, [renamed])

    assert restored.tables[0].sheet_name == "Крепёж новый"
    assert restored.tables[0].default_category_path == "Крепёж новый"


def test_recipe_keeps_an_explicit_category_that_differs_from_the_sheet_name():
    grid = _grid(HEADERS, ROWS_A, name="Крепёж")
    decision = build_auto_mapping([grid])
    decision.tables[0].default_category_path = "Каталог / Метизы"

    recipe = recipe_from_decision(decision, [grid])

    assert recipe["sheets"][0].get("default_category_source") is None
    assert recipe["sheets"][0]["default_category_path"] == "Каталог / Метизы"
    restored = recipe_to_decision(recipe, [grid])
    assert restored.tables[0].default_category_path == "Каталог / Метизы"


def test_legacy_recipe_derives_safe_sheet_category_context():
    meaningful = _grid(HEADERS, ROWS_A, name="Крепёж")
    truncated = _grid(HEADERS, ROWS_A, name="Каталог - Оптические кабели свя")
    recipe = recipe_from_decision(build_auto_mapping([meaningful]), [meaningful])
    recipe["sheets"][0].pop("default_category_path")

    restored = recipe_to_decision(recipe, [meaningful])
    assert restored.tables[0].default_category_path == "Крепёж"

    recipe["sheets"][0]["match"]["name_pattern"] = truncated.name
    restored_truncated = recipe_to_decision(recipe, [truncated])
    assert restored_truncated.tables[0].default_category_path is None


def test_tabular_rule_rejects_non_scalar_attribute_kind() -> None:
    with pytest.raises(HTTPException) as exc_info:
        _validate_recipe_format(
            {
                "sheets": [
                    {
                        "columns": [
                            {"index": 0, "role": "attribute", "attribute_value_kind": "jsonb"}
                        ]
                    }
                ]
            }
        )

    assert exc_info.value.status_code == 422
    assert "jsonb" in str(exc_info.value.detail)


def test_recipe_does_not_report_empty_worksheet_as_non_price():
    price_sheet = _grid(HEADERS, ROWS_A)
    empty_sheet = SheetGrid(name="Пустой лист")
    decision = build_auto_mapping([price_sheet, empty_sheet])

    restored = recipe_to_decision(
        recipe_from_decision(decision, [price_sheet, empty_sheet]),
        [price_sheet, empty_sheet],
    )

    assert restored.auto_ok
    assert restored.non_price_sheets == []


def test_recipe_prefill_aligns_by_header_name():
    grid = _grid(HEADERS, ROWS_A)
    decision = build_auto_mapping([grid])
    recipe = recipe_from_decision(decision, [grid])

    # Same headers, but the article column moved to the end.
    moved = _grid(["Наименование", "Ед. изм.", "Цена", "Артикул"], [["Муфта", "шт", 5, "A-1"]])
    prefill = recipe_prefill_mapping(recipe, [moved])
    cols = {c["role"]: c["index"] for c in prefill["tables"][0]["columns"]}
    assert cols["article"] == 3
    assert cols["name"] == 0
    assert cols["price"] == 2


def test_stale_rule_reports_ignored_description_and_technical_attribute_columns():
    grid = _grid(TARGET_HEADERS, TARGET_ROWS, name="Прайс")
    stale_recipe = {
        "sheets": [
            {
                "match": {"name_pattern": "Прайс", "or_index": 0},
                "skip": False,
                "header_row": 0,
                "columns": [
                    {"index": 0, "role": "name", "header": "наименование"},
                    {"index": 1, "role": "article", "header": "артикул"},
                    {"index": 2, "role": "ignore", "header": "толщина металла"},
                    {
                        "index": 3,
                        "role": "dimensions",
                        "header": "габаритные размеры длина глубина высота мм",
                    },
                    {"index": 4, "role": "ignore", "header": "описание"},
                    {
                        "index": 5,
                        "role": "price",
                        "header": "цена с ндс",
                        "price_descriptor": {"kind": "other", "vat_mode": "included"},
                    },
                ],
            }
        ]
    }

    gaps = recipe_mapping_gaps(stale_recipe, [grid])

    assert {gap["expected_role"] for gap in gaps} >= {"description", "attribute"}


def test_stale_rule_reports_factory_price_mapped_as_manufacturer():
    """A legacy ``завод`` alias must not keep a real price as manufacturer."""
    grid = _grid(
        ["Наименование", "Новая заводская цена 12.01.2026", "Розничная цена"],
        [["Шкаф пожарный", "4 500", "5 000"]],
        name="Прайс",
    )
    stale_recipe = {
        "sheets": [
            {
                "match": {"name_pattern": "Прайс", "or_index": 0},
                "skip": False,
                "header_row": 0,
                "columns": [
                    {"index": 0, "role": "name", "header": "наименование"},
                    {
                        "index": 1,
                        "role": "manufacturer",
                        "header": "новая заводская цена 12 01 2026",
                    },
                    {
                        "index": 2,
                        "role": "price",
                        "header": "розничная цена",
                        "price_descriptor": {"kind": "retail", "vat_mode": "unknown"},
                    },
                ],
            }
        ]
    }

    assert recipe_mapping_gaps(stale_recipe, [grid]) == [
        {
            "sheet": "Прайс",
            "index": 1,
            "header": "Новая заводская цена 12.01.2026",
            "expected_role": "price",
            "actual_role": "manufacturer",
        }
    ]


def test_recipe_roundtrip_keeps_vertical_tables_in_their_own_columns():
    """Saved rules retain section bounds for vertically stacked price tables."""
    grid = SheetGrid(
        name="Прайс",
        rows=[
            [
                Cell(value="Наименование"),
                Cell(value="Цвет"),
                Cell(value="Толщина металла"),
                Cell(value="Цена"),
            ],
            [
                Cell(value="Шкаф"),
                Cell(value="серый"),
                Cell(value="0,8 мм"),
                Cell(value="100"),
            ],
            [Cell(value=None), Cell(value=None), Cell(value=None), Cell(value=None)],
            [
                Cell(value="Кол-во секций"),
                Cell(value="Толщина металла"),
                Cell(value="Опт-1"),
                Cell(value=None),
            ],
            [
                Cell(value="2"),
                Cell(value="0,8 мм"),
                Cell(value="200"),
                Cell(value=None),
            ],
        ],
    )

    decision = build_auto_mapping([grid])
    assert len(decision.tables) == 2

    restored = recipe_to_decision(recipe_from_decision(decision, [grid]), [grid])
    result = parse_grids([grid], restored)

    assert [item.source_name for item in result.items] == ["Шкаф", "2"]
    assert [item.values[0].amount for item in result.items] == [
        Decimal("100.00"),
        Decimal("200.00"),
    ]
    assert [
        value.raw_value
        for item in result.items
        for value in item.attribute_values
        if value.source_header == "Толщина металла"
    ] == ["0,8 мм", "0,8 мм"]


# --- DB-level: draft rule creation and exact detection (sqlite) --------------

def _make_xlsx(headers: list[str], rows: list[list[object]]) -> bytes:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Прайс"
    ws.append(headers)
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


async def _setup_version(session, supplier: Supplier, filename: str, version_no: int):
    pf = PriceFile(
        supplier_id=supplier.id,
        original_filename=filename,
        storage_key=f"sha256/x/{filename}",
        sha256=f"hash{version_no}",
        size_bytes=1,
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    session.add(pf)
    await session.flush()
    version = PriceVersion(
        supplier_id=supplier.id,
        price_file_id=pf.id,
        version_no=version_no,
        status=VersionStatus.parsing.value,
    )
    session.add(version)
    await session.flush()
    return version, pf


def test_auto_draft_rule_then_exact_match():
    asyncio.run(_run_auto_draft_rule_then_exact_match())


async def _run_auto_draft_rule_then_exact_match():
    engine = create_async_engine("sqlite+aiosqlite://")
    session_factory = async_sessionmaker(engine, expire_on_commit=False)
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)

    data_v1 = _make_xlsx(HEADERS, ROWS_A)
    data_v2 = _make_xlsx(HEADERS, ROWS_B)  # same structure, different rows

    async with session_factory() as session:
        supplier = Supplier(code="t-1", name="Test")
        session.add(supplier)
        await session.flush()

        v1, pf1 = await _setup_version(session, supplier, "price_v1.xlsx", 1)
        status = await parse_and_persist(session, v1, pf1, data_v1)
        await session.commit()

        assert status == VersionStatus.validated.value
        assert v1.loader == "xlsx"
        assert v1.rule_match == "none"

        rules = (await session.execute(select(ParseRule))).scalars().all()
        assert len(rules) == 1
        rule = rules[0]
        assert rule.status == ParseRuleStatus.draft.value
        assert rule.provenance == ParseRuleProvenance.auto.value
        assert rule.created_from_version_id == v1.id
        assert rule.validation and rule.validation["passed"] is True
        assert v1.parse_rule_id == rule.id

        # Second file with the same structure -> exact rule match, no new rule.
        v2, pf2 = await _setup_version(session, supplier, "price_v2.xlsx", 2)
        status2 = await parse_and_persist(session, v2, pf2, data_v2)
        await session.commit()

        assert status2 == VersionStatus.validated.value
        assert v2.rule_match == "exact"
        assert v2.parse_rule_id == rule.id
        rules_after = (await session.execute(select(ParseRule))).scalars().all()
        assert len(rules_after) == 1

        items = (
            (
                await session.execute(
                    select(SupplierPriceItem).where(SupplierPriceItem.price_version_id == v2.id)
                )
            )
            .scalars()
            .all()
        )
        assert len(items) == len(ROWS_B)
    await engine.dispose()


def test_stale_auto_rule_with_ignored_recognized_column_is_rebuilt():
    asyncio.run(_run_stale_auto_rule_with_ignored_recognized_column_is_rebuilt())


async def _run_stale_auto_rule_with_ignored_recognized_column_is_rebuilt():
    engine = create_async_engine("sqlite+aiosqlite://")
    session_factory = async_sessionmaker(engine, expire_on_commit=False)
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)

    data = _make_xlsx(COMPACT_DIMENSIONS_HEADERS, COMPACT_DIMENSIONS_ROWS)
    grid = _grid(COMPACT_DIMENSIONS_HEADERS, COMPACT_DIMENSIONS_ROWS, name="Прайс")
    fingerprint = compute_fingerprint("xlsx", [grid])
    stale_recipe = {
        "sheets": [
            {
                "match": {"name_pattern": "Прайс", "or_index": 0},
                "skip": False,
                "header_row": 0,
                "columns": [
                    {"index": 0, "role": "name", "header": "наименование"},
                    {"index": 1, "role": "ignore", "header": "шир выс гл"},
                    {
                        "index": 2,
                        "role": "price",
                        "header": "цена",
                        "price_descriptor": {"kind": "other", "vat_mode": "unknown"},
                    },
                ],
            }
        ],
        "options": {"currency": "RUB"},
    }

    async with session_factory() as session:
        supplier = Supplier(code="t-stale", name="Stale rule")
        session.add(supplier)
        await session.flush()
        stale_rule = ParseRule(
            supplier_id=supplier.id,
            name="старое правило",
            loader="xlsx",
            header_hash=compute_header_hash(fingerprint),
            fingerprint=fingerprint,
            recipe=stale_recipe,
            revision=1,
            status=ParseRuleStatus.active.value,
            provenance=ParseRuleProvenance.auto.value,
        )
        session.add(stale_rule)
        await session.flush()

        version, price_file = await _setup_version(session, supplier, "compact.xlsx", 1)
        status = await parse_and_persist(session, version, price_file, data)
        await session.commit()

        assert status == VersionStatus.validated.value
        assert stale_rule.status == ParseRuleStatus.deprecated.value
        assert version.rule_match == "none"
        replacement = (
            await session.execute(
                select(ParseRule).where(ParseRule.id != stale_rule.id)
            )
        ).scalars().one()
        assert replacement.revision == 2
        roles = {column["role"] for column in replacement.recipe["sheets"][0]["columns"]}
        assert "dimensions" in roles
    await engine.dispose()


def test_maintenance_reparse_bypasses_stale_auto_partial_rule():
    asyncio.run(_run_maintenance_reparse_bypasses_stale_auto_partial_rule())


async def _run_maintenance_reparse_bypasses_stale_auto_partial_rule():
    """A stale auto rule must not block an allow_rule_mutations=False run."""
    engine = create_async_engine("sqlite+aiosqlite://")
    session_factory = async_sessionmaker(engine, expire_on_commit=False)
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)

    old_headers = ["Наименование", "Цена"]
    current_headers = ["Наименование", "Цена", "Материал корпуса"]
    old_grid = _grid(old_headers, [["Шкаф", "100"]], name="Прайс")
    old_decision = build_auto_mapping([old_grid])
    old_fingerprint = compute_fingerprint("xlsx", [old_grid])

    async with session_factory() as session:
        supplier = Supplier(code="t-maintenance-stale", name="Maintenance stale")
        session.add(supplier)
        await session.flush()
        stale_rule = ParseRule(
            supplier_id=supplier.id,
            name="устаревший авторазбор",
            loader="xlsx",
            header_hash=compute_header_hash(old_fingerprint),
            fingerprint=old_fingerprint,
            recipe=recipe_from_decision(old_decision, [old_grid]),
            revision=1,
            status=ParseRuleStatus.active.value,
            provenance=ParseRuleProvenance.auto.value,
        )
        session.add(stale_rule)
        await session.flush()

        version, price_file = await _setup_version(session, supplier, "new-columns.xlsx", 1)
        status = await parse_and_persist(
            session,
            version,
            price_file,
            _make_xlsx(current_headers, [["Шкаф", "100", "Сталь"]]),
            allow_rule_mutations=False,
        )
        await session.commit()

        assert status == VersionStatus.validated.value
        assert version.rule_match == RuleMatch.none.value
        assert stale_rule.status == ParseRuleStatus.active.value
        item = (
            await session.execute(
                select(SupplierPriceItem).where(SupplierPriceItem.price_version_id == version.id)
            )
        ).scalars().one()
        assert [value.raw_value for value in item.attribute_values] == ["Сталь"]
    await engine.dispose()


def test_dimensions_are_persisted_with_the_item():
    asyncio.run(_run_dimensions_are_persisted_with_the_item())


async def _run_dimensions_are_persisted_with_the_item():
    engine = create_async_engine("sqlite+aiosqlite://")
    session_factory = async_sessionmaker(engine, expire_on_commit=False)
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)

    async with session_factory() as session:
        supplier = Supplier(code="t-dimensions", name="Dimensions")
        session.add(supplier)
        await session.flush()
        version, price_file = await _setup_version(session, supplier, "dimensions.xlsx", 1)

        status = await parse_and_persist(
            session,
            version,
            price_file,
            _make_xlsx(
                ["Наименование", "Габариты", "Цена"],
                [["Шкаф", "540х640х230", 1188]],
            ),
        )
        await session.commit()

        assert status == VersionStatus.validated.value
        item = (
            await session.execute(
                select(SupplierPriceItem).where(SupplierPriceItem.price_version_id == version.id)
            )
        ).scalars().one()
        assert getattr(item, "dimensions", None) == "540х640х230"
    await engine.dispose()


def test_description_and_technical_attributes_are_persisted_with_the_item():
    asyncio.run(_run_description_and_technical_attributes_are_persisted_with_the_item())


async def _run_description_and_technical_attributes_are_persisted_with_the_item():
    engine = create_async_engine("sqlite+aiosqlite://")
    session_factory = async_sessionmaker(engine, expire_on_commit=False)
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)

    async with session_factory() as session:
        supplier = Supplier(code="t-description", name="Description")
        session.add(supplier)
        await session.flush()
        version, price_file = await _setup_version(session, supplier, "target.xlsx", 1)

        status = await parse_and_persist(
            session,
            version,
            price_file,
            _make_xlsx(TARGET_HEADERS, TARGET_ROWS),
        )
        await session.commit()

        assert status == VersionStatus.validated.value
        item = (
            await session.execute(
                select(SupplierPriceItem).where(SupplierPriceItem.price_version_id == version.id)
            )
        ).scalars().one()
        assert item.description == "Передвижной пост для средств пожаротушения"
        assert [(value.raw_value, value.value_text) for value in item.attribute_values] == [
            ("Лист 0.8 мм", "Лист 0.8 мм")
        ]
    await engine.dispose()


def test_items_api_exposes_description_and_technical_attributes():
    asyncio.run(_run_items_api_exposes_description_and_technical_attributes())


async def _run_items_api_exposes_description_and_technical_attributes():
    engine = create_async_engine("sqlite+aiosqlite://")
    session_factory = async_sessionmaker(engine, expire_on_commit=False)
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)

    async with session_factory() as session:
        supplier = Supplier(code="t-item-api", name="Item API")
        session.add(supplier)
        await session.flush()
        version, price_file = await _setup_version(session, supplier, "target.xlsx", 1)
        await parse_and_persist(
            session,
            version,
            price_file,
            _make_xlsx(TARGET_HEADERS, TARGET_ROWS),
        )
        await session.commit()

        page = await list_items(version.id, query=None, limit=50, offset=0, session=session)

        assert page.total == 1
        assert page.items[0].description == "Передвижной пост для средств пожаротушения"
        assert [attribute.model_dump(exclude={"id", "parse_status"}) for attribute in page.items[0].attributes] == [
            {
                "code": page.items[0].attributes[0].code,
                "name": "Толщина металла",
                "value_kind": "text",
                "value": "Лист 0.8 мм",
                "raw_value": "Лист 0.8 мм",
                "unit": None,
            }
        ]
    await engine.dispose()


def test_json_export_exposes_description_and_technical_attributes():
    asyncio.run(_run_json_export_exposes_description_and_technical_attributes())


async def _run_json_export_exposes_description_and_technical_attributes():
    engine = create_async_engine("sqlite+aiosqlite://")
    session_factory = async_sessionmaker(engine, expire_on_commit=False)
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)

    async with session_factory() as session:
        supplier = Supplier(code="t-export", name="Export")
        session.add(supplier)
        await session.flush()
        version, price_file = await _setup_version(session, supplier, "target.xlsx", 1)
        await parse_and_persist(
            session,
            version,
            price_file,
            _make_xlsx(TARGET_HEADERS, TARGET_ROWS),
        )
        await session.commit()

        response = await _export_json(version.id, session)
        item = json.loads(response.body)["items"][0]

        assert item["description"] == "Передвижной пост для средств пожаротушения"
        assert item["attributes"][0]["name"] == "Толщина металла"
        assert item["attributes"][0]["value"] == "Лист 0.8 мм"
    await engine.dispose()


def test_json_export_serializes_category_ids():
    asyncio.run(_run_json_export_serializes_category_ids())


async def _run_json_export_serializes_category_ids():
    engine = create_async_engine("sqlite+aiosqlite://")
    session_factory = async_sessionmaker(engine, expire_on_commit=False)
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)

    async with session_factory() as session:
        supplier = Supplier(code="t-export-category", name="Export category")
        session.add(supplier)
        await session.flush()
        version, price_file = await _setup_version(session, supplier, "target.xlsx", 1)
        await parse_and_persist(
            session,
            version,
            price_file,
            _make_xlsx(TARGET_HEADERS, TARGET_ROWS),
        )
        supplier_category = SupplierCategory(
            supplier_id=supplier.id,
            source_name="Огнетушители",
            normalized_name="огнетушители",
            source_path="Огнетушители",
            normalized_path="огнетушители",
        )
        global_category = GlobalCategory(
            name="Пожарное оборудование",
            normalized_name="пожарное оборудование",
            path="Пожарное оборудование",
            normalized_path="пожарное оборудование",
        )
        session.add_all([supplier_category, global_category])
        await session.flush()
        session.add(
            SupplierCategoryMapping(
                supplier_category_id=supplier_category.id,
                global_category_id=global_category.id,
                mapping_source="manual",
                created_by="test",
            )
        )
        item = await session.scalar(
            select(SupplierPriceItem).where(
                SupplierPriceItem.price_version_id == version.id
            )
        )
        assert item is not None
        item.category_path = supplier_category.source_path
        item.supplier_category_id = supplier_category.id
        await session.commit()
        version_id = version.id
        supplier_category_id = supplier_category.id
        global_category_id = global_category.id
        session.expire_all()

        response = await _export_json(version_id, session)
        category = json.loads(response.body)["items"][0]["category"]

        assert category["supplier_category_id"] == str(supplier_category_id)
        assert category["global_category_id"] == str(global_category_id)
    await engine.dispose()


def test_partial_match_prefills_mapping():
    asyncio.run(_run_partial_match_prefills_mapping())


async def _run_partial_match_prefills_mapping():
    engine = create_async_engine("sqlite+aiosqlite://")
    session_factory = async_sessionmaker(engine, expire_on_commit=False)
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)

    async with session_factory() as session:
        supplier = Supplier(code="t-2", name="Test2")
        session.add(supplier)
        await session.flush()

        v1, pf1 = await _setup_version(session, supplier, "base.xlsx", 1)
        await parse_and_persist(session, v1, pf1, _make_xlsx(HEADERS, ROWS_A))
        await session.commit()

        # Structure changed: one header renamed -> no exact hash, partial >= 0.6.
        changed = _make_xlsx(
            ["Артикул", "Наименование", "Ед. изм.", "Розничная цена"], ROWS_A
        )
        v2, pf2 = await _setup_version(session, supplier, "changed.xlsx", 2)
        status = await parse_and_persist(session, v2, pf2, changed)
        await session.commit()

        assert status == VersionStatus.needs_mapping.value
        assert v2.rule_match == "partial"
        assert v2.parse_rule_id is not None
        assert v2.mapping and v2.mapping["tables"], "маппинг должен быть предзаполнен"
        roles = {c["role"] for c in v2.mapping["tables"][0]["columns"]}
        assert "name" in roles
    await engine.dispose()


def test_user_mapping_upserts_active_rule():
    asyncio.run(_run_user_mapping_upserts_active_rule())


def test_failed_auto_mapping_stores_rule_generation_prompt():
    asyncio.run(_run_failed_auto_mapping_stores_rule_generation_prompt())


async def _run_failed_auto_mapping_stores_rule_generation_prompt():
    engine = create_async_engine("sqlite+aiosqlite://")
    session_factory = async_sessionmaker(engine, expire_on_commit=False)
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)

    data = _make_xlsx(
        ["Наименование", "Изображение"],
        [["Скамья СК001", "[image]"]],
    )

    async with session_factory() as session:
        supplier = Supplier(code="t-4", name="Test4")
        session.add(supplier)
        await session.flush()
        version, price_file = await _setup_version(session, supplier, "unique.xlsx", 1)

        status = await parse_and_persist(session, version, price_file, data)

        assert status == VersionStatus.needs_mapping.value
        assert version.rule_generation
        assert version.rule_generation["model"] == "gpt-4o"
        assert version.rule_generation["status"] == "prepared"
        assert "Наименование" in version.rule_generation["prompt"]
    await engine.dispose()


async def _run_user_mapping_upserts_active_rule():
    engine = create_async_engine("sqlite+aiosqlite://")
    session_factory = async_sessionmaker(engine, expire_on_commit=False)
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)

    async with session_factory() as session:
        supplier = Supplier(code="t-3", name="Test3")
        session.add(supplier)
        await session.flush()

        v1, pf1 = await _setup_version(session, supplier, "user.xlsx", 1)
        v1.mapping = {
            "tables": [
                {
                    "sheet_name": "Прайс",
                    "columns": [
                        {"index": 0, "role": "article"},
                        {"index": 1, "role": "name"},
                        {"index": 3, "role": "price",
                         "price_descriptor": {"kind": "retail", "vat_mode": "included"}},
                    ],
                }
            ]
        }
        status = await parse_and_persist(session, v1, pf1, _make_xlsx(HEADERS, ROWS_A))
        await session.commit()

        assert status == VersionStatus.validated.value
        rules = (await session.execute(select(ParseRule))).scalars().all()
        assert len(rules) == 1
        assert rules[0].status == ParseRuleStatus.active.value
        assert rules[0].provenance == ParseRuleProvenance.user_confirmed.value
        assert v1.parse_rule_id == rules[0].id
    await engine.dispose()


def test_auto_draft_recipe_reconstructs_from_source_file(monkeypatch):
    asyncio.run(_run_auto_draft_recipe_reconstructs_from_source_file(monkeypatch))


async def _run_auto_draft_recipe_reconstructs_from_source_file(monkeypatch) -> None:
    engine = create_async_engine("sqlite+aiosqlite://")
    session_factory = async_sessionmaker(engine, expire_on_commit=False)
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)

    data = _make_xlsx(HEADERS, ROWS_A)

    async with session_factory() as session:
        project = Project(
            id=ORIGINAL_PROJECT_ID,
            slug="original-prices",
            name="Оригинальные прайсы",
            is_default=True,
        )
        supplier = Supplier(code="t-auto-draft", name="Auto draft")
        session.add_all([project, supplier])
        await session.flush()
        version, price_file = await _setup_version(session, supplier, "auto-draft.xlsx", 1)
        version.status = VersionStatus.validated.value
        rule = ParseRule(
            supplier_id=supplier.id,
            name="Правило с пустым рецептом",
            loader="xlsx",
            header_hash="d" * 64,
            fingerprint={},
            recipe={"sheets": []},
            revision=3,
            status="draft",
            provenance="manual",
            created_from_version_id=version.id,
        )
        session.add(rule)
        await session.flush()
        monkeypatch.setattr(storage, "get_price_file", lambda _key: data)

        outcome = await auto_draft_recipe(rule.id, project, session)

        assert outcome.sheet_count == 1
        assert outcome.passed is True
        assert outcome.metrics["rows_accepted"] == len(ROWS_A)
        columns = outcome.recipe["sheets"][0]["columns"]
        roles = {column["role"] for column in columns}
        assert "name" in roles
        assert "price" in roles
        # The stored rule is left untouched; only the draft payload is returned.
        assert rule.recipe == {"sheets": []}
    await engine.dispose()


def test_cloned_rule_keeps_source_version_for_complete_auto_reconstruction():
    asyncio.run(_run_cloned_rule_keeps_source_version_for_complete_auto_reconstruction())


async def _run_cloned_rule_keeps_source_version_for_complete_auto_reconstruction() -> None:
    engine = create_async_engine("sqlite+aiosqlite://")
    session_factory = async_sessionmaker(engine, expire_on_commit=False)
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)
    try:
        async with session_factory() as session:
            project = Project(
                id=ORIGINAL_PROJECT_ID,
                slug="original-prices",
                name="Оригинальные прайсы",
                is_default=True,
            )
            supplier = Supplier(code="clone-source", name="Clone source")
            session.add_all([project, supplier])
            await session.flush()
            version, _price_file = await _setup_version(session, supplier, "source.xlsx", 1)
            source = ParseRule(
                supplier_id=supplier.id,
                name="Архивное правило",
                loader="xlsx",
                header_hash="f" * 64,
                fingerprint={},
                recipe={"sheets": []},
                revision=1,
                status=ParseRuleStatus.deprecated.value,
                provenance=ParseRuleProvenance.manual.value,
                created_from_version_id=version.id,
            )
            session.add(source)
            await session.flush()

            clone = await clone_rule_to_draft(source.id, project, session)

            assert clone.created_from_version_id == version.id
    finally:
        await engine.dispose()


def test_validation_and_activation_reject_source_item_loss(monkeypatch):
    asyncio.run(_run_validation_and_activation_reject_source_item_loss(monkeypatch))


async def _run_validation_and_activation_reject_source_item_loss(monkeypatch) -> None:
    engine = create_async_engine("sqlite+aiosqlite://")
    session_factory = async_sessionmaker(engine, expire_on_commit=False)
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)
    data = _make_xlsx(HEADERS, ROWS_A)
    try:
        async with session_factory() as session:
            project = Project(
                id=ORIGINAL_PROJECT_ID,
                slug="original-prices",
                name="Оригинальные прайсы",
                is_default=True,
            )
            supplier = Supplier(code="source-loss", name="Source loss")
            session.add_all([project, supplier])
            await session.flush()
            version, _price_file = await _setup_version(
                session, supplier, "source-loss.xlsx", 1
            )
            version.item_count = 3
            rule = ParseRule(
                supplier_id=supplier.id,
                name="Неполное правило",
                loader="xlsx",
                header_hash="f" * 64,
                fingerprint={},
                recipe=recipe_from_decision(
                    build_auto_mapping([_grid(HEADERS, ROWS_A)]),
                    [_grid(HEADERS, ROWS_A)],
                ),
                revision=1,
                status=ParseRuleStatus.draft.value,
                provenance=ParseRuleProvenance.manual.value,
                created_from_version_id=version.id,
            )
            session.add(rule)
            await session.flush()
            monkeypatch.setattr(storage, "get_price_file", lambda _key: data)

            validation = await validate_rule(
                rule.id,
                RuleValidateIn(price_version_id=version.id),
                project,
                session,
            )

            assert validation.passed is False
            assert validation.metrics["source_items"] == 3
            assert validation.metrics["items_delta"] == -1
            assert validation.samples[0].name == "Муфта 1"
            assert validation.samples[0].article == "A-1"
            assert validation.samples[0].source_sheet == "Прайс"
            assert validation.samples[0].source_row == 2
            assert validation.samples[0].prices[0].amount == "100.50"
            assert validation.samples[0].prices[0].source_column_index == 3
            assert rule.validation["samples"][0]["name"] == "Муфта 1"

            # Legacy validations may already have been stored as passed. The
            # activation endpoint must independently enforce the same guard.
            rule.validation = {
                "passed": True,
                "price_version_id": str(version.id),
                "metrics": {"items_with_price": 2},
            }
            with pytest.raises(HTTPException) as exc_info:
                await set_rule_status(rule.id, ParseRuleStatusUpdate(status="active"), False, project, session)
            assert exc_info.value.status_code == 409
            assert "2 из 3" in exc_info.value.detail
            assert rule.status == ParseRuleStatus.draft.value
    finally:
        await engine.dispose()


def test_rule_validation_and_matrix_tools_reject_another_supplier_version():
    asyncio.run(_run_rule_validation_and_matrix_tools_reject_another_supplier_version())


async def _run_rule_validation_and_matrix_tools_reject_another_supplier_version() -> None:
    engine = create_async_engine("sqlite+aiosqlite://")
    session_factory = async_sessionmaker(engine, expire_on_commit=False)
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)
    try:
        async with session_factory() as session:
            project = Project(
                id=ORIGINAL_PROJECT_ID,
                slug="original-prices",
                name="Оригинальные прайсы",
                is_default=True,
            )
            rule_supplier = Supplier(code="rule-supplier", name="Rule supplier")
            other_supplier = Supplier(code="other-supplier", name="Other supplier")
            session.add_all([project, rule_supplier, other_supplier])
            await session.flush()
            version, _price_file = await _setup_version(
                session, other_supplier, "other.xlsx", 1
            )
            recipe = {
                "format": "workbook_layout_v1",
                "profile": "price_new_v1",
                "hydrant_matrix": {},
                "service_matrix": {},
            }
            rule = ParseRule(
                supplier_id=rule_supplier.id,
                name="Supplier-scoped rule",
                loader="xlsx",
                header_hash="e" * 64,
                fingerprint={},
                recipe=recipe,
                revision=1,
                status=ParseRuleStatus.draft.value,
                provenance=ParseRuleProvenance.manual.value,
            )
            session.add(rule)
            await session.flush()

            calls = [
                lambda: validate_rule(
                    rule.id,
                    RuleValidateIn(price_version_id=version.id),
                    project,
                    session,
                ),
                lambda: preview_rule_matrix(
                    rule.id,
                    RuleMatrixPreviewIn(
                        price_version_id=version.id,
                        source_section_id="matrix-1",
                        recipe=recipe,
                    ),
                    project,
                    session,
                ),
                lambda: rebind_rule_matrix(
                    rule.id,
                    RuleMatrixRebindIn(
                        price_version_id=version.id,
                        recipe=recipe,
                    ),
                    project,
                    session,
                ),
            ]
            for call in calls:
                with pytest.raises(HTTPException) as exc_info:
                    await call()
                assert exc_info.value.status_code == 409
                assert "другому поставщику" in exc_info.value.detail

            rule.validation = {
                "passed": True,
                "price_version_id": str(version.id),
                "metrics": {"items_with_price": 1},
            }
            with pytest.raises(HTTPException) as exc_info:
                await set_rule_status(rule.id, ParseRuleStatusUpdate(status="active"), False, project, session)
            assert exc_info.value.status_code == 409
            assert "другому поставщику" in exc_info.value.detail
            assert rule.status == ParseRuleStatus.draft.value
    finally:
        await engine.dispose()


@requires_price_new
def test_auto_draft_requires_model_generation_for_matrix_workbook(monkeypatch):
    asyncio.run(_run_auto_draft_requires_model_generation_for_matrix_workbook(monkeypatch))


async def _run_auto_draft_requires_model_generation_for_matrix_workbook(monkeypatch) -> None:
    engine = create_async_engine("sqlite+aiosqlite://")
    session_factory = async_sessionmaker(engine, expire_on_commit=False)
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)

    data = PRICE_NEW_FILE.read_bytes()
    try:
        async with session_factory() as session:
            project = Project(
                id=ORIGINAL_PROJECT_ID,
                slug="original-prices",
                name="Оригинальные прайсы",
                is_default=True,
            )
            supplier = Supplier(code="price-new-layout", name="Supplier F layout")
            session.add_all([project, supplier])
            await session.flush()
            version, _price_file = await _setup_version(
                session, supplier, PRICE_NEW_FILE.name, 1
            )
            version.status = VersionStatus.validated.value
            rule = ParseRule(
                supplier_id=supplier.id,
                name="Восстановить полное правило",
                loader="xlsx",
                header_hash="e" * 64,
                fingerprint={},
                recipe={"sheets": []},
                revision=1,
                status="draft",
                provenance="manual",
                created_from_version_id=version.id,
            )
            session.add(rule)
            await session.flush()
            monkeypatch.setattr(storage, "get_price_file", lambda _key: data)

            with pytest.raises(HTTPException) as exc_info:
                await auto_draft_recipe(rule.id, project, session)

            assert exc_info.value.status_code == 409
            assert "Сгенерировать правило из прайса" in exc_info.value.detail
    finally:
        await engine.dispose()


@requires_price_new
def test_matrix_import_waits_for_generated_rule_instead_of_builtin_recipe():
    asyncio.run(_run_matrix_import_waits_for_generated_rule_instead_of_builtin_recipe())


async def _run_matrix_import_waits_for_generated_rule_instead_of_builtin_recipe() -> None:
    engine = create_async_engine("sqlite+aiosqlite://")
    session_factory = async_sessionmaker(engine, expire_on_commit=False)
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)

    data = PRICE_NEW_FILE.read_bytes()
    try:
        async with session_factory() as session:
            project = Project(
                id=ORIGINAL_PROJECT_ID,
                slug="original-prices",
                name="Оригинальные прайсы",
                is_default=True,
            )
            supplier = Supplier(code="price-new-auto-layout", name="Supplier F auto layout")
            session.add_all([project, supplier])
            await session.flush()

            first, first_file = await _setup_version(
                session, supplier, PRICE_NEW_FILE.name, 1
            )
            first_status = await parse_and_persist(session, first, first_file, data)
            await session.commit()

            rules = (await session.execute(select(ParseRule))).scalars().all()
            assert first_status == VersionStatus.needs_mapping.value
            assert first.item_count == 0
            assert rules == []
            assert first.parse_rule_id is None
            assert first.rule_generation["mode"] == "section_proposals"
            assert first.rule_generation["matrix_sections_count"] >= 1

            second, second_file = await _setup_version(
                session, supplier, "renamed-price-new.xlsx", 2
            )
            second_status = await parse_and_persist(session, second, second_file, data)
            await session.commit()

            assert second_status == VersionStatus.needs_mapping.value
            assert second.item_count == 0
            assert second.parse_rule_id is None
            assert second.rule_match == RuleMatch.none.value
            rules_after = (await session.execute(select(ParseRule))).scalars().all()
            assert rules_after == []
    finally:
        await engine.dispose()
