# -*- coding: utf-8 -*-
"""Regression tests for the Supplier E validation fix.

Parser glues quantity-tier "ladder" rows (merged-fill product name, unit
column holding "31-60"/"61-100"/">100") into the parent position. The review
gates used to count those absorbed rows as missed (35 in dry-run, 12 on the
import-level completeness gate), and a fresh review opened with an empty
mapping so columns 5/6/7 were lost. These tests pin the fix:
1. ParsedItem.glued_source_rows records the absorbed Excel rows.
2. dry_run_summary treats glued rows as covered (missed == 0).
3. reconcile_version passes when persisted items carry glued_source_rows.
4. A generic sheet-columns recipe prefills the review mapping, keeping the
   unit and second price pair columns.
"""
from __future__ import annotations

import asyncio
import io

import openpyxl
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

from app.models import (
    Base,
    PriceFile,
    PriceVersion,
    Supplier,
    SupplierPriceItem,
)
from app.parsing.grid import Cell, SheetGrid
from app.parsing.mapping import ColumnMapping, MappingDecision, TableMapping
from app.parsing.pipeline import parse_grids
from app.reconciliation import reconcile_version
from app.review_validation import dry_run_summary

SHEET_NAME = "\u041f\u0420\u0410\u0419\u0421 \u0434\u0438\u043b.\u0440\u043e\u0437"
H_PROD = "\u043f\u0440\u043e\u0434\u0443\u043a\u0446\u0438\u044f"
H_DEALER = "\u0446\u0435\u043d\u0430 \u0434\u043b\u044f \u0434\u0438\u043b\u0435\u0440\u043e\u0432 \u0440\u0443\u0431"
H_UNIT = "\u0435\u0434 \u0438\u0437\u043c"
H_RETAIL = "\u0446\u0435\u043d\u0430 \u0440\u043e\u0437\u043d\u0438\u0447\u043d\u0430\u044f \u0440\u0443\u0431"
NAME = "\u0411\u043e\u043b\u0442 \u041c8"


def _cell_row(*values):
    return [Cell(value=value) for value in values]


def _merged(value):
    return Cell(value=value, is_merged_fill=True)


def _tekhno_grid() -> SheetGrid:
    """Miniature of the Supplier E sheet: product + three quantity-tier ladders.

    Grid row 0 is the header, row 1 the product, rows 2-4 the ladders whose
    name is a merged fill of the product name (Excel rows 3-5).
    """
    header = _cell_row(
        "", H_PROD, H_PROD, H_DEALER, H_DEALER, H_UNIT, H_RETAIL, H_RETAIL
    )
    product = _cell_row("", NAME, NAME, 100, 110, "\u0448\u0442", 120, 130)
    ladders = []
    for label, low_price, high_price in [("31-60", 110, 120), ("61-100", 105, 115), (">100", 100, 110)]:
        ladders.append(
            [
                Cell(value=""),
                _merged(NAME),
                _merged(NAME),
                Cell(value=""),
                Cell(value=""),
                Cell(value=label),
                Cell(value=low_price),
                Cell(value=high_price),
            ]
        )
    return SheetGrid(name=SHEET_NAME, rows=[header, product, *ladders])


def _tekhno_decision() -> MappingDecision:
    table = TableMapping(
        sheet_name=SHEET_NAME,
        header_row_index=0,
        columns=[
            ColumnMapping(index=0, role="ignore", header=""),
            ColumnMapping(index=1, role="attribute", header=H_PROD),
            ColumnMapping(index=2, role="name", header=H_PROD),
            ColumnMapping(index=3, role="price", header=H_DEALER,
                          price_descriptor={"kind": "dealer", "vat_mode": "excluded"}),
            ColumnMapping(index=4, role="price", header=H_DEALER,
                          price_descriptor={"kind": "dealer", "vat_mode": "included"}),
            ColumnMapping(index=5, role="unit", header=H_UNIT),
            ColumnMapping(index=6, role="price", header=H_RETAIL,
                          price_descriptor={"kind": "retail", "vat_mode": "excluded"}),
            ColumnMapping(index=7, role="price", header=H_RETAIL,
                          price_descriptor={"kind": "retail", "vat_mode": "included"}),
        ],
    )
    return MappingDecision(tables=[table], auto_ok=True)


def _tekhno_recipe() -> dict:
    columns = [
        {"index": 0, "role": "ignore"},
        {"index": 1, "role": "attribute", "header": H_PROD},
        {"index": 2, "role": "name", "header": H_PROD},
        {"index": 3, "role": "price", "header": H_DEALER,
         "price_descriptor": {"kind": "dealer", "vat_mode": "excluded"}},
        {"index": 4, "role": "price", "header": H_DEALER,
         "price_descriptor": {"kind": "dealer", "vat_mode": "included"}},
        {"index": 5, "role": "unit", "header": H_UNIT},
        {"index": 6, "role": "price", "header": H_RETAIL,
         "price_descriptor": {"kind": "retail", "vat_mode": "excluded"}},
        {"index": 7, "role": "price", "header": H_RETAIL,
         "price_descriptor": {"kind": "retail", "vat_mode": "included"}},
    ]
    return {
        "sheets": [
            {
                "match": {"name_pattern": SHEET_NAME, "or_index": 0},
                "skip": False,
                "header_row": 0,
                "columns": columns,
            }
        ],
        "options": {},
    }


def test_ladder_rows_are_glued_into_parent_and_recorded() -> None:
    result = parse_grids([_tekhno_grid()], _tekhno_decision())
    assert result.errors_count == 0
    # Three ladders must not become three extra catalogue positions.
    assert len(result.items) == 1
    item = result.items[0]
    assert item.source_row == 2
    assert sorted(item.glued_source_rows) == [3, 4, 5]
    assert len(item.values) == 10


def test_dry_run_counts_glued_ladder_rows_as_covered() -> None:
    grid = _tekhno_grid()
    result = parse_grids([grid], _tekhno_decision())
    summary = dry_run_summary(
        result=result,
        grids=[grid],
        decisions=[],
        currency_override=None,
        current_item_count=0,
    )
    assert summary["metrics"]["confirmed_missed_positions"] == 0, summary["metrics"]["missed_coordinates"]
    assert summary["passed"] is True


def test_reconcile_passes_when_item_records_glued_rows() -> None:
    asyncio.run(_reconcile_case())


async def _reconcile_case() -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = SHEET_NAME
    sheet.append(["", H_PROD, H_PROD, H_DEALER, H_DEALER, H_UNIT, H_RETAIL, H_RETAIL])
    sheet.append(["", NAME, NAME, 100, 110, "\u0448\u0442", 120, 130])
    # Ladders reuse the product name through a vertical merge (B3:C5).
    sheet.merge_cells(start_row=3, start_column=2, end_row=5, end_column=3)
    sheet.cell(row=3, column=6, value="31-60")
    sheet.cell(row=4, column=6, value="61-100")
    sheet.cell(row=5, column=6, value=">100")
    sheet.cell(row=3, column=7, value=110)
    sheet.cell(row=3, column=8, value=120)
    sheet.cell(row=4, column=7, value=105)
    sheet.cell(row=4, column=8, value=115)
    sheet.cell(row=5, column=7, value=100)
    sheet.cell(row=5, column=8, value=110)
    output = io.BytesIO()
    workbook.save(output)
    data = output.getvalue()

    engine = create_async_engine("sqlite+aiosqlite://")
    factory = async_sessionmaker(engine, expire_on_commit=False)
    async with engine.begin() as connection:
        await connection.run_sync(Base.metadata.create_all)
    try:
        async with factory() as session:
            supplier = Supplier(code="tekhno-fix", name="T")
            session.add(supplier)
            await session.flush()
            price_file = PriceFile(
                supplier_id=supplier.id,
                original_filename="tekhno-fix.xlsx",
                storage_key="tekhno-fix-file",
                sha256="f" * 64,
                size_bytes=len(data),
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
            session.add(price_file)
            await session.flush()
            version = PriceVersion(
                supplier_id=supplier.id, price_file_id=price_file.id, version_no=1
            )
            session.add(version)
            await session.flush()
            session.add(
                SupplierPriceItem(
                    price_version_id=version.id,
                    source_name=NAME,
                    normalized_name="\u0431\u043e\u043b\u0442 \u043c8",
                    source_sheet=SHEET_NAME,
                    source_row=2,
                    source_panel_index=0,
                    source_column_start=0,
                    source_column_end=8,
                    glued_source_rows=[3, 4, 5],
                )
            )
            await session.commit()
            summary = await reconcile_version(
                session, version=version, data=data, filename="tekhno-fix.xlsx"
            )
            assert summary["passed"] is True
            assert summary["confirmed_missed_positions"] == 0, summary["missed_coordinates"]
            assert summary["missed_coordinates"] == []
    finally:
        await engine.dispose()


def test_review_mapping_prefill_keeps_unit_and_second_price_pair() -> None:
    from app.review_service import _recipe_review_mapping

    mapping = _recipe_review_mapping(_tekhno_recipe(), [_tekhno_grid()])
    assert mapping is not None
    assert len(mapping["tables"]) == 1
    table = mapping["tables"][0]
    assert table["sheet_name"] == SHEET_NAME
    assert table["header_row_index"] == 0
    by_index = {column["index"]: column for column in table["columns"]}
    # Every column of the sheet is explicit, ignores included.
    assert sorted(by_index) == list(range(8))
    assert by_index[0]["role"] == "ignore"
    assert by_index[2]["role"] == "name"
    assert by_index[5]["role"] == "unit"
    assert by_index[3]["role"] == "price"
    assert by_index[3]["price_descriptor"] == {"kind": "dealer", "vat_mode": "excluded"}
    assert by_index[6]["role"] == "price"
    assert by_index[6]["price_descriptor"] == {"kind": "retail", "vat_mode": "excluded"}
    assert by_index[7]["role"] == "price"
    assert by_index[7]["price_descriptor"] == {"kind": "retail", "vat_mode": "included"}